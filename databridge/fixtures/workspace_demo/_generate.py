"""Regenerate the Larkspur Analytics synthetic demo corpus under ``files/``.

Every company, person, address, email, phone number, IBAN, card number and key in this corpus is
INVENTED. Domains use the reserved ``.example`` TLD; phone numbers use the UK drama range
(020 7946 0xxx); the IBANs are the published documentation examples; the card is the well-known
Luhn-valid test number; the private key body is not a key.

Run from anywhere:  python fixtures/workspace_demo/_generate.py
Needs: reportlab (PDF) and python-docx (DOCX). Output is deterministic apart from PDF/DOCX
internal timestamps.
"""
from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).resolve().parent
FILES = HERE / "files"

# ── plain-text documents (.md / .txt) ────────────────────────────────────────────────────────────
TEXT: dict[str, str] = {}

TEXT["projects/route-optimisation/brief.md"] = """\
# Project brief — Route Optimisation Pilot (PRJ-101)

**Customer:** Pellucid Freightways Ltd (customer id C-1001)
**Contract:** MSA-2026-014
**Project manager:** Tobias Reinholt (Larkspur Analytics, Head of Delivery)
**Customer sponsor:** Rowan Achterberg, Head of Network Planning, Pellucid Freightways
**Start:** 2026-07-01 · **Planned end:** 2026-12-18

## Goal
Reduce empty running on Pellucid's southern trunk network by re-sequencing depot departures
from the Harbourgate depot in Southampton. The pilot covers 42 vehicles and three depots.

## Scope
1. Ingest twelve months of telematics exports.
2. Build a baseline of empty miles per lane.
3. Deliver a weekly re-sequencing recommendation to the planning desk.

## Team
- Tobias Reinholt — delivery lead
- Ines Marchetti-Holm — modelling (part-time, shared with PRJ-102)
- Rowan Achterberg — customer sponsor

## Related
Kickoff notes: meetings/2026-07-14_kickoff_pellucid.txt. Invoicing under MSA-2026-014.
"""

TEXT["projects/route-optimisation/notes.md"] = """\
# Working notes — Route Optimisation Pilot

2026-07-22 — Telematics export from Pellucid arrived (Rowan Achterberg sent it). 11 of 12
months present; August 2025 missing. Tobias Reinholt to chase.

2026-08-05 — Baseline built for the Harbourgate depot. Ines Marchetti-Holm flagged that two
lanes share a vehicle pool, so the empty-mile split between them is an assumption, not a
measurement. Recorded as an open question for the steering group.

2026-09-02 — First re-sequencing recommendation issued to the Pellucid planning desk.
"""

TEXT["projects/demand-forecasting/brief.md"] = """\
# Project brief — Demand Forecasting Rollout (PRJ-102)

**Customer:** Corriedale Grocers Co-operative (customer id C-1002)
**Contract:** SOW-2026-021 (under the Corriedale framework agreement)
**Project manager:** Ines Marchetti-Holm (Larkspur Analytics, Lead Data Scientist)
**Customer sponsor:** Hamish Oduya, Supply Chain Director, Corriedale Grocers
**Start:** 2026-08-01 · **Planned end:** 2027-01-29

## Goal
Replace the spreadsheet forecast for fresh produce across 18 Corriedale stores in Leeds and
Harrogate with a weekly store-by-SKU forecast.

## Team
- Ines Marchetti-Holm — project manager and modelling lead
- Petra Lindqvist-Shaw — Corriedale category analyst, owns the SKU master
- Callum Driscoll-Park — Larkspur finance, invoicing

## Related
Steering notes: meetings/2026-08-19_steering_corriedale.docx
"""

TEXT["projects/demand-forecasting/notes.md"] = """\
# Working notes — Demand Forecasting Rollout

2026-08-07 — SKU master received from Petra Lindqvist-Shaw. 2,140 active SKUs.

2026-08-19 — Steering meeting with Hamish Oduya. Agreed to start with the Leeds stores only.

2026-09-10 — Backtest complete for Leeds. Hand-off of the weekly forecast to Corriedale planning
is scheduled for the week of 2026-10-05.
"""

TEXT["meetings/2026-07-14_kickoff_pellucid.txt"] = """\
MEETING NOTES — Kickoff, Route Optimisation Pilot (PRJ-101)
Date: 2026-07-14   Location: Pellucid Freightways, Harbourgate depot, Southampton

Attendees
  Larkspur Analytics: Tobias Reinholt, Ines Marchetti-Holm
  Pellucid Freightways: Rowan Achterberg, Selin Kaya-Morrow (Procurement)

Discussion
  - Rowan Achterberg walked through the southern trunk network and the three pilot depots.
  - Selin Kaya-Morrow confirmed invoices go to Pellucid accounts payable, quoting MSA-2026-014.
  - Contact for data questions: rowan.achterberg@pellucid-freightways.example, +44 20 7946 0321.

Actions
  1. Rowan Achterberg to send twelve months of telematics exports by 2026-07-21.
  2. Tobias Reinholt to circulate the baseline method by 2026-08-01.
"""

TEXT["hr/Compensation_Bands_2026.md"] = """\
STRICTLY CONFIDENTIAL

# Compensation bands 2026 — Larkspur Analytics

Owner: Ada Nwachukwu-Barr (People & Culture)

| Band | Title | Salary range |
|---|---|---|
| L1 | Analyst | GBP 34,000 – 41,000 |
| L2 | Data Scientist | GBP 45,000 – 58,000 |
| L3 | Lead Data Scientist | GBP 62,000 – 76,000 |
| L4 | Head of function | GBP 82,000 – 98,000 |

The 2026 bonus pool is set at 6% of eligible salaries.
"""

TEXT["engineering/deploy_notes.txt"] = """\
Deploy notes — forecasting service (internal)
Author: Ines Marchetti-Holm

The staging box is reached through the bastion. Someone pasted the deploy key below while
debugging; it must be removed and rotated. (Planted for the screening demo: this is NOT a real
key — the body is filler text — but the header alone must block the file.)

-----BEGIN OPENSSH PRIVATE KEY-----
RkFLRS1LRVktRk9SLVRIRS1XQUxOVVQtU0NSRUVOSU5HLURFTU8tTk9ULUEtUkVBTC1L
RVktTk9ULUEtUkVBTC1LRVktTk9ULUEtUkVBTC1LRVktTk9ULUEtUkVBTC1LRVktRkFL
RS1GQUtFLUZBS0UtRkFLRS1GQUtFLUZBS0UtRkFLRS1GQUtFLUZBS0UtRkFLRS1GQUtF
-----END OPENSSH PRIVATE KEY-----

Service restarts cleanly after the config change on 2026-09-03.
"""

TEXT["sales/Rate_Card_Draft_v2.md"] = """\
# Rate card — draft v2 (not yet approved)

Prepared by Mira Oyelaran for the 2027 renewals.

## Pricing

| Service | Unit | Rate |
|---|---|---|
| Discovery sprint | per sprint | £14,500 |
| Forecasting model build | per model | £38,000 |
| Route optimisation pilot | per depot per month | £4,200 |
| Managed service | per month | £6,750 |

Discounts above 10% need sign-off from Mira Oyelaran.
"""

TEXT["planning/Q3 #roadmap/roadmap.md"] = """\
# Q3 roadmap (folder name contains '#': a collector that URL-addresses folders sees zero files here)

Owner: Mira Oyelaran (CEO)

- PRJ-101 Route Optimisation Pilot (Pellucid Freightways) — on track; Tobias Reinholt.
- PRJ-102 Demand Forecasting Rollout (Corriedale Grocers Co-operative) — Leeds first; Ines Marchetti-Holm.
- PRJ-103 Leak Detection Discovery (Tamsinbay Water Authority) — proposal stage. Contact:
  Oluwaseun Pettifer, Asset Data Manager. Discovery planned for 2026-11.
- Hire: one L2 data scientist in Bristol.
"""

# ── binary documents (.pdf / .docx) — paragraphs, rendered below ──────────────────────────────
PDFS: dict[str, list[str]] = {
    "contracts/Pellucid/MSA-2026-014_Pellucid_Freightways.pdf": [
        "MASTER SERVICES AGREEMENT  MSA-2026-014",
        "This Master Services Agreement is made on 2026-06-24 between Larkspur Analytics Ltd, "
        "Unit 4 Carrow Mill, Thornbury Lane, Bristol (the Supplier) and Pellucid Freightways Ltd, "
        "Harbourgate, Southampton (the Customer).",
        "1. Services. The Supplier will deliver the services described in each statement of work, "
        "the first being the Route Optimisation Pilot (PRJ-101).",
        "2. Fees. Fees are set out in each statement of work and invoiced monthly in arrears.",
        "3. Confidentiality. Each party shall keep the other party's confidential information "
        "secret and use it only to perform this agreement.",
        "4. Liability. Each party's total liability is limited to GBP 5,000,000, and each party "
        "shall maintain professional indemnity insurance of GBP 10,000,000.",
        "5. Term. This agreement runs from 2026-07-01 for 24 months.",
        "Signed for the Supplier: Mira Oyelaran, Chief Executive, Larkspur Analytics Ltd.",
        "Signed for the Customer: Selin Kaya-Morrow, Head of Procurement, Pellucid Freightways Ltd.",
    ],
    "finance/invoices/INV-2026-0041.pdf": [
        "INVOICE  INV-2026-0041",
        "From: Larkspur Analytics Ltd, Unit 4 Carrow Mill, Thornbury Lane, Bristol",
        "Bill to: Pellucid Freightways Ltd, Accounts Payable, Harbourgate, Southampton",
        "Customer id: C-1001   Project: PRJ-101 Route Optimisation Pilot   Contract: MSA-2026-014",
        "Issue date: 2026-08-31   Due date: 2026-09-30",
        "Route optimisation pilot, August 2026 (3 depots): GBP 18,500.00",
        "VAT at 20%: GBP 3,700.00      Total due: GBP 22,200.00",
        "Pay by bank transfer to IBAN GB82 WEST 1234 5698 7654 32 quoting INV-2026-0041.",
        "Queries: callum.driscoll-park@larkspur-analytics.example or +44 20 7946 0188.",
        "Issued by Callum Driscoll-Park, Finance Manager.",
    ],
    "finance/invoices/INV-2026-0057.pdf": [
        "INVOICE  INV-2026-0057",
        "From: Larkspur Analytics Ltd, Unit 4 Carrow Mill, Thornbury Lane, Bristol",
        "Bill to: Corriedale Grocers Co-operative, Finance Office, Leeds",
        "Customer id: C-1002   Project: PRJ-102 Demand Forecasting Rollout   Contract: SOW-2026-021",
        "Issue date: 2026-09-15   Due date: 2026-10-15",
        "Demand forecasting rollout, milestone 1 (Leeds backtest): GBP 12,750.00",
        "Sensor calibration kit, part no. DE44 LSPR 2026 0457 11, supplied at no charge.",
        "VAT at 20%: GBP 2,550.00      Total due: GBP 15,300.00",
        "Pay by bank transfer to IBAN GB29 NWBK 6016 1331 9268 19 quoting INV-2026-0057.",
        "Deposit already taken on card 4111 1111 1111 1111 (test card) on 2026-09-01.",
        "Queries: callum.driscoll-park@larkspur-analytics.example or +44 20 7946 0188.",
        "Issued by Callum Driscoll-Park, Finance Manager.",
    ],
}

DOCXS: dict[str, list[str]] = {
    "contracts/Corriedale/SOW-2026-021_Demand_Forecasting.docx": [
        "STATEMENT OF WORK  SOW-2026-021",
        "Between Larkspur Analytics Ltd and Corriedale Grocers Co-operative, Leeds.",
        "Project: PRJ-102 Demand Forecasting Rollout. Start 2026-08-01; end 2027-01-29.",
        "Larkspur project manager: Ines Marchetti-Holm. Corriedale sponsor: Hamish Oduya, "
        "Supply Chain Director.",
        "Deliverables: a weekly store-by-SKU forecast for 18 stores; a hand-over pack; training "
        "for the Corriedale planning team led by Petra Lindqvist-Shaw.",
        "Total contract value: GBP 84,000, invoiced against four milestones.",
        "Signed: Mira Oyelaran (Larkspur Analytics) and Hamish Oduya (Corriedale Grocers).",
    ],
    "meetings/2026-08-19_steering_corriedale.docx": [
        "Steering meeting — Demand Forecasting Rollout (PRJ-102)",
        "Date: 2026-08-19. Location: Corriedale Grocers head office, Leeds.",
        "Attendees: Hamish Oduya, Petra Lindqvist-Shaw (Corriedale Grocers); Ines Marchetti-Holm, "
        "Callum Driscoll-Park (Larkspur Analytics).",
        "Decision: start with the Leeds stores only; Harrogate follows once the Leeds backtest "
        "passes.",
        "Action: Petra Lindqvist-Shaw to freeze the SKU master by 2026-08-29.",
        "Action: Callum Driscoll-Park to raise the milestone 1 invoice after the backtest.",
    ],
    "hr/Remote_Working_Policy.docx": [
        "Remote Working Policy — Larkspur Analytics",
        "Owner: Ada Nwachukwu-Barr, People & Culture. Version 3, effective 2026-04-01.",
        "Staff may work remotely up to three days a week by agreement with their manager.",
        "Requests to change working pattern are handled in a strictly confidential manner and "
        "are seen only by the employee, their manager and People & Culture.",
        "Customer data must only be accessed through company-managed laptops and never stored "
        "on personal devices.",
        "Questions about this policy go to Ada Nwachukwu-Barr.",
    ],
}


def _write_pdf(path: Path, paras: list[str]) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=A4, title=path.stem,
                            author="Larkspur Analytics (synthetic)")
    story = []
    for i, p in enumerate(paras):
        story.append(Paragraph(p, styles["Heading2"] if i == 0 else styles["BodyText"]))
        story.append(Spacer(1, 6))
    doc.build(story)


def _write_docx(path: Path, paras: list[str]) -> None:
    import docx

    d = docx.Document()
    d.core_properties.author = "Larkspur Analytics (synthetic)"
    d.add_heading(paras[0], level=1)
    for p in paras[1:]:
        d.add_paragraph(p)
    d.save(str(path))


def main() -> None:
    for rel, body in TEXT.items():
        p = FILES / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(body, encoding="utf-8", newline="\n")
    for rel, paras in PDFS.items():
        p = FILES / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        _write_pdf(p, paras)
    for rel, paras in DOCXS.items():
        p = FILES / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        _write_docx(p, paras)
    print(f"wrote {len(TEXT) + len(PDFS) + len(DOCXS)} files under {FILES}")


if __name__ == "__main__":
    main()
