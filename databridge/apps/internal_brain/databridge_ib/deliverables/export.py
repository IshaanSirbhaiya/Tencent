"""Exports, built on the server from the STORED, CHECKED content only (never from anything a
browser sends). CSV and .pptx (P1 Slides) for tables and charts. File names are built from fixed
parts only.

The formula guard (OWASP CSV injection; Opus review #14): a cell whose first non-space character
is ``= + - @``, a tab, CR or LF, or a full-width ``＝ ＋ － ＠`` (also after leading spaces,
no-break spaces or other Unicode spaces) gets a leading ``'``, except a plain number such as
``-0.5%``. It applies to every cell, the Sources block included.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import datetime
from typing import Any, Mapping, Sequence


CSV_CAP = 256 * 1024
_TRIGGER = re.compile(r"^[\s\u00a0\u1680\u2000-\u200b\u202f\u205f\u3000\ufeff]*[=+\-@\t\r\n\uff1d\uff0b\uff0d\uff20]")
_PLAIN_NUMBER = re.compile(r"^-?\d[\d,]*(?:\.\d+)?%?$")
FILENAME_RE = re.compile(r"^walnut-[a-z]+-[0-9a-f]{8}-\d{8}\.(?:pptx|csv)$")
CONTENT_TYPES = {"csv": "text/csv; charset=utf-8",
                 "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation"}


def safe_link(url: Any) -> str:
    """Only an http(s) link from our own citation metadata; anything else is left out."""
    u = str(url or "")
    return u if re.match(r"^https?://[^\s\"<>]+$", u) else ""


def guard(cell: Any) -> str:
    s = "" if cell is None else str(cell)
    if _TRIGGER.match(s) and not _PLAIN_NUMBER.match(s.strip()):
        return "'" + s
    return s


def filename(kind: str, req_id: str, when: datetime, ext: str) -> str:
    """``walnut-<kind>-<req_id[:8]>-<yyyymmdd>.<ext>``; a request id that is not hex gives
    ``00000000`` (never an error, never the id itself)."""
    head = (req_id or "")[:8].lower()
    hex8 = head if re.fullmatch(r"[0-9a-f]{8}", head) else "00000000"
    name = f"walnut-{kind}-{hex8}-{when:%Y%m%d}.{ext}"
    if not FILENAME_RE.match(name):                 # fixed parts only: this cannot fail
        name = f"walnut-file-00000000-{when:%Y%m%d}.{ext}"
    return name


def content_disposition(name: str) -> str:
    assert FILENAME_RE.match(name)
    return f'attachment; filename="{name}"'


def _markers(tags: Sequence[str], number: Mapping[str, int]) -> str:
    return " ".join(f"[{number[t]}]" for t in dict.fromkeys(tags) if t in number)


def build_csv(kind: str, content: Mapping[str, Any], citations: Sequence[Mapping[str, Any]]) -> bytes:
    number = {c["tag"]: c["n"] for c in citations}
    buf = io.StringIO()
    w = csv.writer(buf, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")

    def row(cells):
        w.writerow([guard(c) for c in cells])

    if kind == "table":
        row(list(content["columns"]) + ["Sources"])
        for r in content["rows"]:
            row([c["text"] for c in r] + [_markers([t for c in r for t in c["tags"]], number)])
    elif kind == "chart":
        row(["Series", "Label", "Value as written", "Value", "Unit", "Source"])
        unit = content.get("currency") or content.get("unit") or ""
        for s in content["series"]:
            for p in s["points"]:
                row([s["name"], p["x"], p["raw"], p["value"], unit, _markers([p["tag"]], number)])
    else:
        raise ValueError("kind")
    w.writerow([])
    row(["Source", "Title", "Platform", "Place", "Version", "Checked at", "Link"])
    seen = set()
    for c in citations:
        if c["doc_id"] in seen:
            continue
        seen.add(c["doc_id"])
        row([f"[{c['n']}]", c.get("title", ""), c.get("platform", ""), c.get("container_name", ""),
             c.get("version", ""), c.get("checked_at", ""), safe_link(c.get("deep_link", ""))])
    return ("﻿" + buf.getvalue()).encode("utf-8")


# ------------------------------------------------------------------------------------------------
# .pptx (P1 "Slides"): built from the SAME stored, re-checked content as the CSV. No model.
# ------------------------------------------------------------------------------------------------
PPTX_CAP = 2 * 1024 * 1024
PPTX_ROWS_PER_SLIDE = 10
#: what the speaker notes say was checked (Opus review #8: never "every number was found")
NOTES_CHECKED = ("Each value on this slide was matched by Walnut's checks to a number written in "
                 "the cited source line, bound to its label. The checks do not prove that a "
                 "comparison between values is meaningful.")
_XML_BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def pptx_available() -> bool:
    """python-pptx (and XlsxWriter, for native charts) importable. Imported lazily, so the IB starts
    without them; without them the deliverable offers no ``pptx`` export."""
    try:
        import pptx  # noqa: F401
        import xlsxwriter  # noqa: F401
    except Exception:                                               # noqa: BLE001
        return False
    return True


def ptext(s: Any) -> str:
    """Every string written into the deck, its chart's embedded workbook included: the CSV formula
    guard (Opus review #4 / T15: python-pptx writes chart categories and series names with
    XlsxWriter ``write()``, which turns a leading ``=`` into a formula), then XML-invalid control
    characters removed (lxml refuses them)."""
    return _XML_BAD.sub("", guard(s))


def build_pptx(kind: str, content: Mapping[str, Any], citations: Sequence[Mapping[str, Any]], *,
               title: str, req_id: str, as_of: datetime) -> bytes:
    """A 16:9 deck: a title slide, the table (native, 10 rows per slide) or the chart (a native,
    editable column or line chart from the checked points), then a Sources slide."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt

    number = {c["tag"]: c["n"] for c in citations}
    head = (req_id or "")[:8].lower()
    rid = head if re.fullmatch(r"[0-9a-f]{8}", head) else "00000000"
    footer = f"Walnut DataBridge · draft for review · request {rid}"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank, title_only, title_layout = prs.slide_layouts[6], prs.slide_layouts[5], prs.slide_layouts[0]

    def notes(slide, tags: Sequence[str]) -> None:
        srcs = [c for c in citations if not tags or c["tag"] in tags]
        lines = [ptext(f"[{c['n']}] {c.get('title', '')} · {c.get('platform', '')} · "
                       f"v{c.get('version', '')}") for c in srcs]
        slide.notes_slide.notes_text_frame.text = "\n".join(lines + [NOTES_CHECKED])

    def foot(slide) -> None:
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(6.95), Inches(12.3), Inches(0.4))
        tb.text_frame.text = ptext(footer)
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(10)

    s = prs.slides.add_slide(title_layout)
    s.shapes.title.text = ptext(title)
    s.placeholders[1].text = ptext(f"Built from Walnut DataBridge request {rid} on "
                                   f"{as_of:%d %b %Y} · Draft for your review")
    notes(s, [])

    if kind == "table":
        cols = list(content["columns"]) + ["Sources"]
        rows = [[c["text"] for c in r] + [_markers([t for c in r for t in c["tags"]], number)]
                for r in content["rows"]]
        for i in range(0, max(1, len(rows)), PPTX_ROWS_PER_SLIDE):
            part = rows[i:i + PPTX_ROWS_PER_SLIDE]
            s = prs.slides.add_slide(title_only)
            s.shapes.title.text = ptext(title if i == 0 else f"{title} (continued)")
            shape = s.shapes.add_table(len(part) + 1, len(cols), Inches(0.5), Inches(1.5),
                                       Inches(12.3), Inches(0.4) * (len(part) + 1))
            tbl = shape.table
            for j, h in enumerate(cols):
                tbl.cell(0, j).text = ptext(h)
            for r_i, r in enumerate(part, 1):
                for j, v in enumerate(r):
                    tbl.cell(r_i, j).text = ptext(v)
                    for run in tbl.cell(r_i, j).text_frame.paragraphs[0].runs:
                        run.font.size = Pt(12)
            notes(s, [t for r in content["rows"][i:i + PPTX_ROWS_PER_SLIDE] for c in r
                      for t in c["tags"]])
            foot(s)
    elif kind == "chart":
        cats: list[str] = []
        for se in content["series"]:
            for p in se["points"]:
                if p["x"] not in cats:
                    cats.append(p["x"])
        data = CategoryChartData(number_format="General")
        data.categories = [ptext(x) for x in cats]
        for se in content["series"]:
            by_x = {p["x"]: p["value"] for p in se["points"]}
            data.add_series(ptext(se["name"]), [by_x.get(x) for x in cats])
        s = prs.slides.add_slide(title_only)
        s.shapes.title.text = ptext(title)
        ctype = (XL_CHART_TYPE.LINE_MARKERS if content.get("type") == "line"
                 else XL_CHART_TYPE.COLUMN_CLUSTERED)
        chart = s.shapes.add_chart(ctype, Inches(0.7), Inches(1.4), Inches(12), Inches(5.3),
                                   data).chart
        chart.has_legend = len(content["series"]) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        plot = chart.plots[0]
        plot.has_data_labels = True
        unit = content.get("currency") or ("%" if content.get("unit") == "percent"
                                           else content.get("unit") or "")
        y = content.get("y_label") or ""
        axis_title = ptext(f"{y} ({unit})" if y and unit and unit != "number" else y or
                           (unit if unit != "number" else ""))
        if axis_title:
            chart.value_axis.has_title = True
            chart.value_axis.axis_title.text_frame.text = axis_title
        if content.get("x_label"):
            chart.category_axis.has_title = True
            chart.category_axis.axis_title.text_frame.text = ptext(content["x_label"])
        notes(s, [p["tag"] for se in content["series"] for p in se["points"]])
        foot(s)
    else:
        raise ValueError("kind")

    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Sources"
    tb = s.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12.3), Inches(5.2))
    tf = tb.text_frame
    tf.word_wrap = True
    seen, first = set(), True
    for c in citations:
        if c["doc_id"] in seen:
            continue
        seen.add(c["doc_id"])
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.text = ptext(f"[{c['n']}] {c.get('title', '')} · {c.get('platform', '')} · "
                          f"{c.get('container_name', '')} · v{c.get('version', '')} · "
                          f"checked {c.get('checked_at', '')}")
        link = safe_link(c.get("deep_link", ""))
        if link:
            run = para.add_run()
            run.text = "  " + link
            run.hyperlink.address = link
    notes(s, [])
    foot(s)

    cp = prs.core_properties
    cp.title = ptext(title)[:255]
    cp.author = "Walnut DataBridge"
    cp.last_modified_by = ""
    cp.comments = ""
    cp.created = cp.modified = as_of.replace(tzinfo=None)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


__all__ = ["guard", "filename", "content_disposition", "build_csv", "build_pptx", "ptext",
           "pptx_available", "CSV_CAP", "PPTX_CAP", "FILENAME_RE", "CONTENT_TYPES", "NOTES_CHECKED"]
