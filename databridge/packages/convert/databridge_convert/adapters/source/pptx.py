"""PPTX SOURCE adapter — bytes → DocumentRecord (C12, §5.2). READ-ONLY.

python-pptx; each slide → a section: the title shape becomes a Heading, remaining text
frames become Paragraph blocks in a top-to-bottom reading-order heuristic. Documented as
lossy (§13.1): animations, transitions, layout geometry, and media are dropped."""
from __future__ import annotations

import io

from pptx import Presentation

from .._docutil import guard_zip_bomb, make_document
from ...ir import Block, DocumentRecord
from .._util import AdapterError


def read(data: bytes) -> DocumentRecord:
    guard_zip_bomb(data, label="pptx")                 # §21.5 — pptx is a zip
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        raise AdapterError(f"pptx: cannot open presentation: {e}") from e

    blocks: list[Block] = []
    for idx, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else ""
        title_id = title_shape.shape_id if title_shape is not None else None
        blocks.append(Block(kind="Heading", level=2, text=title or f"Slide {idx}"))
        # reading-order heuristic: shapes top-to-bottom, then left-to-right (excluding the
        # title, matched by shape_id — shapes.title returns a fresh wrapper each access).
        texted = [s for s in slide.shapes
                  if s.has_text_frame and s.shape_id != title_id]
        for shape in sorted(texted, key=lambda s: (s.top or 0, s.left or 0)):
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip() or para.text.strip()
                if text:
                    blocks.append(Block(kind="Paragraph", text=text))

    loss = ["PPTX is READ-ONLY and lossy (§5.2): animations and transitions dropped",
            "slide layout geometry / positioning simplified to reading order",
            "images, charts, embedded media dropped",
            "speaker notes not included"]
    return make_document(blocks, data, source_id="pptx",
                         source_meta={"slide_count": len(prs.slides)}, loss=loss)
