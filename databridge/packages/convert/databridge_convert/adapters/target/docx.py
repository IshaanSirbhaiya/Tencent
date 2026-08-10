"""DOCX TARGET adapter — DocumentRecord → bytes (C12, §5.2). python-docx writes headings,
paragraphs, and tables (the §13.4 DOCX report path builds on this)."""
from __future__ import annotations

import io

import docx

from ...ir import DocumentRecord


def write(ir: DocumentRecord) -> bytes:
    document = docx.Document()
    for b in ir.blocks:
        if b.kind == "Heading":
            document.add_heading(b.text, level=max(1, min(b.level or 1, 9)))
        elif b.kind == "Table" and b.rows:
            n_cols = max((len(r) for r in b.rows), default=1)
            table = document.add_table(rows=0, cols=n_cols)
            table.style = "Table Grid"
            for row in b.rows:
                cells = table.add_row().cells
                for j, val in enumerate(row):
                    cells[j].text = str(val)
        elif b.kind == "Code":
            document.add_paragraph(b.text, style="No Spacing")
        elif b.kind == "ListBlock":
            for line in b.text.split("\n"):
                document.add_paragraph(line.lstrip("-*+ ").strip(), style="List Bullet")
        elif b.kind == "PageBreak":
            document.add_page_break()
        else:
            document.add_paragraph(b.text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
