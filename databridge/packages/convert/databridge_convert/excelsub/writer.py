"""Excel writer + Excel reports (09_FINAL_PRD §14.15 / §12.4 / §13.4) — C11.

``WorkbookRecordBuilder`` turns typed columns into styled sheets: header styling, freeze
panes, one RowBatch per sheet, optional TOC, and number formats **derived from semantic
types** (money→currency, timestamp→date). openpyxl is the writer (the §14 library decision).

Two hard rules:
  * **§16.1 semantic-confirmation refusal** — applying semantic-type-driven number formats is
    a plan that depends on inferred semantics, so it REFUSES with ``SEMANTICS_UNCONFIRMED``
    unless the source schema is ``confirmed_by`` a human. Purely structural writes (no
    semantic formatting requested) run freely.
  * **§21.7 formula/CSV injection guard** — any cell whose text starts with ``= + - @`` or a
    tab is apostrophe-prefixed on write so spreadsheets can't execute a smuggled payload.
"""
from __future__ import annotations

import datetime as _dt
import hashlib
import io
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from databridge_core import SemanticType, SourcePointer
from ..ir import ReportIR, RowBatch
from .policies import ExcelRefusal
from .provenance import range_ref

SEMANTICS_UNCONFIRMED = "SEMANTICS_UNCONFIRMED"

# Number formats derived from semantic types (§14.15). Closed map; unknown → General.
NUMBER_FORMATS: dict[SemanticType, str] = {
    SemanticType.AMOUNT_MONEY: '#,##0.00;[Red]-#,##0.00',
    SemanticType.AMOUNT_COUNT: "#,##0",
    SemanticType.SCORE_RATING: "0.00",
    SemanticType.DURATION_SECONDS: "#,##0",
    SemanticType.TIMESTAMP_CREATED: "yyyy-mm-dd hh:mm:ss",
    SemanticType.TIMESTAMP_EVENT: "yyyy-mm-dd hh:mm:ss",
}

_DANGEROUS = ("=", "+", "-", "@", "\t")
_HEADER_FILL = PatternFill("solid", fgColor="2563EB")     # DataBridge blue (08_PRD token)
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def guard_formula(value: Any) -> Any:
    """§21.7: apostrophe-prefix a string whose first non-space char is dangerous so it is
    written as inert TEXT, never an executable formula. Non-strings pass through."""
    if not isinstance(value, str):
        return value
    # Strip only leading SPACES (Excel trims them, so " =cmd" is still a formula) — NOT tabs,
    # since a leading tab is itself a §21.7 trigger and must stay visible to the check.
    stripped = value.lstrip(" ")
    if stripped and stripped[0] in _DANGEROUS:
        return "'" + value
    return value


@dataclass
class SheetSpec:
    """One sheet to write: a name + a RowBatch. ``semantic_types`` (column→SemanticType)
    drives number formats; supplying it REQUIRES a confirmed schema (§16.1)."""

    name: str
    batch: RowBatch
    semantic_types: dict[str, SemanticType] = field(default_factory=dict)


@dataclass
class SheetManifest:
    name: str
    written_range: str
    rows: int
    cols: int
    range_pointer: SourcePointer                      # range-level provenance (§17.1/§14.19)


@dataclass
class WriteResult:
    data: bytes
    sha256: str
    sheets: list[SheetManifest]
    loss_report: list[str] = field(default_factory=list)
    semantic_formats_applied: bool = False


@dataclass
class WriteOptions:
    header_style: bool = True
    freeze_header: bool = True
    toc: bool = False                                 # add a table-of-contents sheet (§14.15)
    number_formats: bool = True                       # apply semantic-type number formats


class WorkbookRecordBuilder:
    """Builds an .xlsx from typed RowBatch sheets (§14.15).

    ``confirmed`` MUST be True for semantic number formats to be applied; otherwise the build
    refuses ``SEMANTICS_UNCONFIRMED`` (§16.1) rather than guessing formats from unconfirmed
    semantics."""

    def __init__(self, options: Optional[WriteOptions] = None) -> None:
        self.options = options or WriteOptions()
        self._specs: list[SheetSpec] = []

    def add_sheet(self, name: str, batch: RowBatch,
                  semantic_types: Optional[dict[str, SemanticType]] = None
                  ) -> "WorkbookRecordBuilder":
        self._specs.append(SheetSpec(name, batch, semantic_types or dict(batch.column_semantics)))
        return self

    def build(self, *, confirmed: bool = False, source_id: str = "workbook") -> WriteResult:
        apply_semantics = self.options.number_formats and any(
            s.semantic_types for s in self._specs)
        if apply_semantics and not confirmed:
            unconfirmed = [s.name for s in self._specs if s.semantic_types]
            raise ExcelRefusal(
                SEMANTICS_UNCONFIRMED,
                f"semantic number formats requested for sheet(s) {unconfirmed} but the source "
                f"schema is not confirmed (§16.1)",
                nearest_alternative="confirm the AnnotatedSchema (human sign-off), then re-run; "
                                    "or write without number_formats for a structural-only export")

        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        loss: list[str] = []
        sheet_manifests: list[SheetManifest] = []
        # determine a source sha for provenance from the first batch's snapshot.
        src_hash = self._specs[0].batch.snapshot.content_hash if self._specs else ""

        if self.options.toc and len(self._specs) > 1:
            self._write_toc(wb)

        for spec in self._specs:
            ws = wb.create_sheet(spec.name)
            cols = spec.batch.columns
            records = spec.batch.records()
            # header
            for c, col in enumerate(cols, start=1):
                cell = ws.cell(row=1, column=c, value=guard_formula(col))
                if self.options.header_style:
                    cell.fill = _HEADER_FILL
                    cell.font = _HEADER_FONT
                    cell.alignment = Alignment(horizontal="center")
            # data
            for r, rec in enumerate(records, start=2):
                for c, col in enumerate(cols, start=1):
                    val = _coerce_for_xlsx(rec.get(col))
                    cell = ws.cell(row=r, column=c, value=guard_formula(val))
                    if apply_semantics:
                        st = spec.semantic_types.get(col)
                        fmt = NUMBER_FORMATS.get(st) if st else None
                        if fmt:
                            cell.number_format = fmt
            if self.options.freeze_header:
                ws.freeze_panes = "A2"
            n_rows, n_cols = len(records), len(cols)
            rng = range_ref(0, 0, n_rows, max(n_cols - 1, 0)) if n_cols else ""
            sheet_manifests.append(SheetManifest(
                name=spec.name, written_range=rng, rows=n_rows, cols=n_cols,
                range_pointer=SourcePointer(
                    source_id=source_id, kind="workbook_cell",
                    workbook_sha=src_hash or "pending", sheet=spec.name,
                    cell_range=rng or "A1:A1")))

        data = _save(wb)
        return WriteResult(
            data=data, sha256=hashlib.sha256(data).hexdigest(), sheets=sheet_manifests,
            loss_report=loss, semantic_formats_applied=apply_semantics)

    def _write_toc(self, wb: openpyxl.Workbook) -> None:
        toc = wb.create_sheet("TOC")
        toc.cell(row=1, column=1, value="Sheet").font = _HEADER_FONT
        toc.cell(row=1, column=2, value="Rows").font = _HEADER_FONT
        toc.cell(row=1, column=1).fill = _HEADER_FILL
        toc.cell(row=1, column=2).fill = _HEADER_FILL
        for i, spec in enumerate(self._specs, start=2):
            toc.cell(row=i, column=1, value=guard_formula(spec.name))   # §21.7 (review 1a)
            toc.cell(row=i, column=2, value=spec.batch.num_rows)
        toc.freeze_panes = "A2"


def _coerce_for_xlsx(v: Any) -> Any:
    """openpyxl accepts str/int/float/bool/datetime/None. Decimal → float (xlsx is float64
    anyway); other exotics → str."""
    if v is None or isinstance(v, (str, int, float, bool, _dt.datetime, _dt.date)):
        return v
    if isinstance(v, Decimal):
        return float(v)
    return str(v)


def _save(wb: openpyxl.Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- single-RowBatch convenience (round-trip / EV-XL write half) ----------------
def write_rowbatch(batch: RowBatch, *, sheet_name: str = "Sheet1") -> bytes:
    """Write one RowBatch to .xlsx bytes (structural; no semantic formats → no confirmation
    needed). Used by the EV-XL write-half round-trip."""
    return WorkbookRecordBuilder(WriteOptions(number_formats=False)).add_sheet(
        sheet_name, batch).build(confirmed=False).data


# --- registry write hook (WorkbookRecord → bytes) -------------------------------
def write_workbook_record(wb_rec: Any) -> bytes:
    """Adapter write entry point: a WorkbookRecord → .xlsx bytes. Each sheet's value grid is
    written with the §21.7 injection guard. Formula/style/chart fidelity is NOT reconstructed
    (values+layout only — recorded as lossy in the matrix)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in wb_rec.sheets:
        ws = wb.create_sheet(sheet.name[:31] or "Sheet")
        for r, row in enumerate(sheet.cells, start=1):
            for c, val in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=guard_formula(_coerce_for_xlsx(val)))
    if not wb_rec.sheets:
        wb.create_sheet("Sheet1")
    return _save(wb)


# --- ReportIR → workbook (§12.4 / §13.4) ----------------------------------------
def render_report(report: ReportIR, *, title: str = "Report") -> bytes:
    """Render a ReportIR to an .xlsx report (§12.4). Deterministic: headings as bold rows,
    paragraphs as text rows, Table blocks as cell ranges. Claims are not re-verified here
    (that is C17's job); the renderer only lays out approved content."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]
    r = 1
    for block in report.blocks:
        if block.kind == "Heading":
            cell = ws.cell(row=r, column=1, value=guard_formula(block.text))
            cell.font = Font(bold=True, size=14 - min((block.level or 1), 4))
            r += 1
        elif block.kind == "Paragraph":
            ws.cell(row=r, column=1, value=guard_formula(block.text))
            r += 1
        elif block.kind == "Table" and block.rows:
            for tr in block.rows:
                for c, val in enumerate(tr, start=1):
                    ws.cell(row=r, column=c, value=guard_formula(val))
                r += 1
        else:
            ws.cell(row=r, column=1, value=guard_formula(block.text))
            r += 1
        r += 1                                        # blank spacer row between blocks
    return _save(wb)
