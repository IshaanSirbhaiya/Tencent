"""C11 — Excel writer + Excel reports (09_FINAL_PRD §14.15 / §12.4 / §21.7).

Covers: the WorkbookRecordBuilder (typed columns, semantic number formats, freeze panes,
TOC), the §16.1 SEMANTICS_UNCONFIRMED refusal, the §21.7 formula-injection guard (an
EV-SEC canary), ReportIR→workbook rendering, and the EV-XL write-half round-trip."""
from __future__ import annotations

import datetime as _dt

import openpyxl
import pytest

from databridge_core import SemanticType, SnapshotIdentity, SourcePointer
from databridge_convert.excelsub import (
    ExcelRefusal,
    ReadOptions,
    WorkbookRecordBuilder,
    WriteOptions,
    guard_formula,
    read_workbook,
    render_report,
    write_rowbatch,
)
from databridge_convert.excelsub.writer import NUMBER_FORMATS, SEMANTICS_UNCONFIRMED
from databridge_convert.generate import extract_facts
from databridge_convert.ir import Block, ReportIR, RowBatch
from databridge_convert.render import ReportBuilder

SNAP = SnapshotIdentity(source_id="s", as_of=_dt.datetime(2026, 1, 1, tzinfo=_dt.timezone.utc),
                        content_hash="abc123", schema_hash="sh")


# --- §21.7 formula/CSV injection guard (EV-SEC canary) --------------------------
INJECTION_PAYLOADS = ["=cmd|' /c calc'!A1", "+SUM(1)", "-2+3", "@SUM(A1)", "\tTAB",
                      "=HYPERLINK(\"http://evil\")"]


@pytest.mark.parametrize("payload", INJECTION_PAYLOADS)
def test_formula_injection_guarded(payload):
    guarded = guard_formula(payload)
    assert guarded.startswith("'"), f"{payload!r} not neutralised"


def test_safe_values_not_touched():
    assert guard_formula("hello") == "hello"
    assert guard_formula(42) == 42
    assert guard_formula(3.14) == 3.14


def test_injection_payload_written_inert_and_reads_back_as_text():
    # EV-SEC: a RowBatch carrying a formula-injection payload must be WRITTEN safely.
    rb = RowBatch.from_records([{"note": "=cmd|' /c calc'!A1"}, {"note": "+SUM(A1)"}], SNAP)
    data = write_rowbatch(rb, sheet_name="S")
    # openpyxl must NOT have stored it as a formula cell
    wb = openpyxl.load_workbook(__import__("io").BytesIO(data))
    ws = wb["S"]
    assert ws["A2"].data_type != "f"                  # not a formula
    assert str(ws["A2"].value).startswith("'")        # apostrophe-guarded text
    # reads back as the guarded literal
    back = read_workbook(data, options=ReadOptions(sheets="all")).sheets[0]
    assert back.cells[1][0].startswith("'=cmd")


# --- §16.1 semantic-confirmation refusal ----------------------------------------
def test_unconfirmed_semantics_refused():
    rb = RowBatch.from_records([{"amount": 10.0}], SNAP)
    b = WorkbookRecordBuilder(WriteOptions()).add_sheet("S", rb,
                                                        {"amount": SemanticType.AMOUNT_MONEY})
    with pytest.raises(ExcelRefusal) as ei:
        b.build(confirmed=False)
    assert ei.value.code == SEMANTICS_UNCONFIRMED
    assert ei.value.nearest_alternative


def test_structural_write_no_semantics_runs_without_confirmation():
    rb = RowBatch.from_records([{"a": "1", "b": "x"}], SNAP)
    res = WorkbookRecordBuilder(WriteOptions(number_formats=False)).add_sheet("S", rb).build()
    assert res.semantic_formats_applied is False
    assert res.data[:2] == b"PK"                       # a valid xlsx (zip)


def test_confirmed_semantics_applies_number_formats():
    rb = RowBatch.from_records([{"amount": 1234.5, "created": _dt.datetime(2025, 1, 1)}], SNAP)
    res = (WorkbookRecordBuilder(WriteOptions())
           .add_sheet("S", rb, {"amount": SemanticType.AMOUNT_MONEY,
                                "created": SemanticType.TIMESTAMP_EVENT})
           .build(confirmed=True))
    assert res.semantic_formats_applied
    wb = openpyxl.load_workbook(__import__("io").BytesIO(res.data))
    ws = wb["S"]
    assert ws["A2"].number_format == NUMBER_FORMATS[SemanticType.AMOUNT_MONEY]
    assert ws["B2"].number_format == NUMBER_FORMATS[SemanticType.TIMESTAMP_EVENT]


# --- §14.15 builder features -----------------------------------------------------
def test_freeze_panes_and_header_style():
    rb = RowBatch.from_records([{"a": 1}], SNAP)
    res = WorkbookRecordBuilder(WriteOptions(number_formats=False)).add_sheet("S", rb).build()
    ws = openpyxl.load_workbook(__import__("io").BytesIO(res.data))["S"]
    assert ws.freeze_panes == "A2"
    assert ws["A1"].font.bold is True


def test_toc_sheet_added_for_multisheet():
    rb = RowBatch.from_records([{"a": 1}], SNAP)
    res = (WorkbookRecordBuilder(WriteOptions(toc=True, number_formats=False))
           .add_sheet("One", rb).add_sheet("Two", rb).build())
    names = [s.name for s in read_workbook(res.data, options=ReadOptions(sheets="all")).sheets]
    assert names[0] == "TOC"


def test_range_provenance_in_manifest():
    rb = RowBatch.from_records([{"a": 1, "b": 2}, {"a": 3, "b": 4}], SNAP)
    res = WorkbookRecordBuilder(WriteOptions(number_formats=False)).add_sheet("S", rb).build()
    m = res.sheets[0]
    assert m.written_range == "A1:B3"                  # header + 2 rows × 2 cols
    assert m.range_pointer.kind == "workbook_cell"
    assert m.range_pointer.cell_range == "A1:B3"


# --- §12.4 ReportIR → workbook ---------------------------------------------------
def _real_report_ir() -> ReportIR:
    """A REAL ReportIR produced by the C17→C18 path (replacing the Prompt-10 hand-built fixture):
    C17 ``extract_facts`` renders the canonical money fact; C18 ``ReportBuilder``/``finalise_report``
    builds the blocks + frozen, cited Claim. The Excel writer (§12.4) then lays it out."""
    ptr = SourcePointer(source_id="wb", kind="workbook_cell", workbook_sha="x", sheet="S", cell="B2")
    facts = extract_facts({"merchant": "Acme", "rev": 100.0},
                          kinds={"rev": "money"}, pointers={"rev": ptr})
    rev = facts[0]                                       # canonical rendered "$100.00"
    b = ReportBuilder()
    b.add_heading("Quarterly", 1)
    b.add_grounded_paragraph(f"Revenue rose to {rev.rendered}.", [rev.pointer])
    b.add_table([["merchant", "rev"], ["Acme", rev.rendered]])
    return b.finalise(SNAP)


def test_report_render_lays_out_blocks():
    rep = _real_report_ir()                             # C17→C18, not a hand-built fixture
    assert rep.claims and all(c.fact_pointers for c in rep.claims)   # real, cited claims
    data = render_report(rep)
    wb = read_workbook(data, options=ReadOptions(sheets="all")).sheets[0]
    flat = [c for row in wb.cells for c in row if c is not None]
    assert "Quarterly" in flat and "Revenue rose to $100.00." in flat
    assert "merchant" in flat and "Acme" in flat


def test_report_render_guards_injection():
    rep = ReportIR(snapshot=SNAP, blocks=[Block(kind="Paragraph", text="=cmd|'/c calc'!A1")])
    data = render_report(rep)
    ws = openpyxl.load_workbook(__import__("io").BytesIO(data)).active
    assert ws["A1"].data_type != "f"


# --- EV-XL write half: round-trip a RowBatch through writer → reader -------------
def test_ev_xl_write_half_roundtrip_value_fidelity():
    rows = [
        {"id": "00012345", "qty": 3.0, "amount": 59.97, "when": _dt.datetime(2024, 3, 1, 9, 30),
         "flag": True, "note": "first"},
        {"id": "00000077", "qty": 5.0, "amount": 25.0, "when": _dt.datetime(2024, 3, 2, 10, 0),
         "flag": False, "note": "second"},
    ]
    rb = RowBatch.from_records(rows, SNAP)
    data = write_rowbatch(rb, sheet_name="RT")
    back = read_workbook(data, options=ReadOptions(sheets="all")).sheets[0]
    assert back.cells[0] == ["id", "qty", "amount", "when", "flag", "note"]
    assert back.cells[1] == ["00012345", 3.0, 59.97, _dt.datetime(2024, 3, 1, 9, 30), True, "first"]
    assert back.cells[2] == ["00000077", 5.0, 25.0, _dt.datetime(2024, 3, 2, 10, 0), False, "second"]
