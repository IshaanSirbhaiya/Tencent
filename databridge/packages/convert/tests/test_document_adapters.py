"""C12 — document adapters + container mode (09_FINAL_PRD §5.2/§11.1/§13.1).

Per-adapter parsing/serialisation, container-mode round-trips (PDF/DOCX/PPTX/HTML →
Markdown), MANDATORY loss reports (empty loss on a lossy pair is itself a bug), and the
structural validator (block/cell counts, char-coverage)."""
from __future__ import annotations

import io

import pytest

from databridge_convert.adapters import documents  # noqa: F401  (registers doc adapters)
from databridge_convert.adapters.container import (
    CHAR_COVERAGE_THRESHOLD,
    convert_document,
    structural_validate,
    to_markdown,
)
from databridge_convert.adapters.source import docx as s_docx
from databridge_convert.adapters.source import html as s_html
from databridge_convert.adapters.source import md as s_md
from databridge_convert.adapters.source import pdf_text as s_pdf
from databridge_convert.adapters.source import pptx as s_pptx
from databridge_convert.adapters.source import txt as s_txt
from databridge_convert.adapters.target import md as t_md
from databridge_convert.router import REGISTRY


# --- fixture builders (binary documents generated deterministically) ------------
def _pdf_bytes() -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.drawString(72, 720, "Invoice 12345")
    c.drawString(72, 700, "Total 99.50")
    c.showPage()
    c.drawString(72, 720, "Second page line")
    c.showPage()
    c.save()
    return buf.getvalue()


def _docx_bytes() -> bytes:
    import docx
    d = docx.Document()
    d.add_heading("Contract", level=1)
    d.add_paragraph("Clause one is binding.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "key"
    t.rows[0].cells[1].text = "value"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarter Review"
    slide.placeholders[1].text = "Revenue grew"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


MD_DOC = b"# Title\n\nFirst paragraph.\n\n- a\n- b\n\n| h1 | h2 |\n| --- | --- |\n| 1 | 2 |\n"
_ARTICLE_P = ("This is the main body content of the page. It is a substantial paragraph of "
              "prose with enough words and sentences that a readability extractor scores it "
              "far above the short navigation and footer links surrounding it. ") * 3
HTML_DOC = (b"<html><body><nav><p>menu nav link</p></nav><article><h1>Headline</h1>"
            + f"<p>{_ARTICLE_P}</p><p>{_ARTICLE_P}</p>".encode()
            + b"<ul><li>one</li><li>two</li></ul>"
            b"<table><tr><th>x</th><th>y</th></tr><tr><td>1</td><td>2</td></tr></table>"
            b"</article><footer><p>footer junk text</p></footer></body></html>")


# --- markdown / txt (lossless, native) ------------------------------------------
def test_markdown_parse_blocks():
    doc = s_md.read(MD_DOC)
    kinds = [b.kind for b in doc.blocks]
    assert kinds == ["Heading", "Paragraph", "ListBlock", "Table"]
    assert doc.blocks[3].rows == [["h1", "h2"], ["1", "2"]]


def test_markdown_roundtrip_identity():
    doc = s_md.read(MD_DOC)
    out = t_md.write(doc)
    doc2 = s_md.read(out)
    assert [b.kind for b in doc2.blocks] == [b.kind for b in doc.blocks]
    assert doc2.blocks[0].text == "Title"


def test_markdown_charindex_populated():
    doc = s_md.read(MD_DOC)
    assert len(doc.char_index.spans) == len(doc.blocks)
    assert doc.char_index.spans[0][0] == 0


def test_txt_roundtrip():
    r = to_markdown(b"Hello there.\n\nAnother paragraph.\n", "txt")
    assert r.validation.char_coverage == 1.0
    assert r.loss_report == []                         # txt→md is lossless


# --- html (lossy: main-content extraction) --------------------------------------
def test_html_main_content_extraction_and_loss():
    r = to_markdown(HTML_DOC, "html")
    text = r.data.decode()
    assert "main body content" in text                 # the article was extracted
    assert "footer junk" not in text                   # boilerplate stripped (readability)
    assert r.loss_report                                # lossy → non-empty


def test_html_full_dom_retains_boilerplate():
    doc = s_html.read(HTML_DOC, full_dom=True)
    text = t_md.write(doc).decode()
    assert "menu nav" in text or "footer junk" in text


# --- pdf (text) ------------------------------------------------------------------
def test_pdf_blocks_pagemap_and_loss():
    doc = s_pdf.read(_pdf_bytes())
    assert any("Invoice 12345" == b.text for b in doc.blocks)
    assert doc.page_map and len(doc.page_map) == 2     # two pages
    assert doc.page_map[0].char_start < doc.page_map[1].char_start
    assert doc.source_meta["loss_report"]


def test_pdf_rejects_non_pdf():
    from databridge_convert.adapters._util import AdapterError
    with pytest.raises(AdapterError):
        s_pdf.read(b"not a pdf")


# --- docx (R+W) ------------------------------------------------------------------
def test_docx_read_blocks_and_table():
    doc = s_docx.read(_docx_bytes())
    assert doc.blocks[0].kind == "Heading" and doc.blocks[0].text == "Contract"
    assert any(b.kind == "Table" for b in doc.blocks)


def test_docx_write_roundtrip():
    from databridge_convert.adapters.target import docx as t_docx
    doc = s_docx.read(_docx_bytes())
    out = t_docx.write(doc)
    doc2 = s_docx.read(out)
    assert doc2.blocks[0].text == "Contract"
    assert any(b.kind == "Table" for b in doc2.blocks)


# --- pptx (R only, lossy) --------------------------------------------------------
def test_pptx_slide_to_section_no_title_dup():
    doc = s_pptx.read(_pptx_bytes())
    assert doc.blocks[0].kind == "Heading" and doc.blocks[0].text == "Quarter Review"
    texts = [b.text for b in doc.blocks]
    assert texts.count("Quarter Review") == 1          # title not duplicated as a paragraph
    assert "Revenue grew" in texts


def test_pptx_is_read_only_in_registry():
    info = REGISTRY.get("pptx")
    assert info.can_read and not info.can_write


# --- container mode: loss reports MANDATORY on lossy pairs (§13.1) ---------------
@pytest.mark.parametrize("builder,fmt", [
    (_pdf_bytes, "pdf"), (_docx_bytes, "docx"), (_pptx_bytes, "pptx"),
])
def test_container_lossy_sources_have_nonempty_loss(builder, fmt):
    r = to_markdown(builder(), fmt)
    assert r.loss_report, f"{fmt}→md produced an EMPTY loss report (§13.1 bug)"


def test_container_lossless_md_has_empty_loss():
    r = convert_document(MD_DOC, "markdown", "markdown")
    assert r.loss_report == []


def test_html_to_md_loss_nonempty():
    r = to_markdown(HTML_DOC, "html")
    assert r.loss_report


# --- structural validator --------------------------------------------------------
def test_structural_validator_char_coverage_and_cells():
    doc = s_docx.read(_docx_bytes())
    out = t_md.write(doc)
    v = structural_validate(doc, out, "markdown")
    assert v.char_coverage >= CHAR_COVERAGE_THRESHOLD
    assert v.output_table_cells == v.source_table_cells   # table cells preserved
    assert v.passed


def test_container_validation_passes_for_docx_to_md():
    r = to_markdown(_docx_bytes(), "docx")
    assert r.validation.passed
    assert r.validation.char_coverage >= CHAR_COVERAGE_THRESHOLD


# --- security: docx/pptx zip-bomb guard (§21.5, review B1) ----------------------
def _zip_bomb() -> bytes:
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("payload.xml", b"A" * (8 * 1024 * 1024))   # 8MB → tiny → huge ratio
    return buf.getvalue()


def test_docx_zip_bomb_refused():
    from databridge_convert.adapters._util import AdapterError
    with pytest.raises(AdapterError):
        s_docx.read(_zip_bomb())


def test_pptx_zip_bomb_refused():
    from databridge_convert.adapters._util import AdapterError
    with pytest.raises(AdapterError):
        s_pptx.read(_zip_bomb())


# --- binary (docx) container target: validator must re-open it (review H2) -------
def test_container_to_docx_target_validates():
    r = convert_document(MD_DOC, "markdown", "docx")
    assert r.data[:2] == b"PK"                          # a real docx (zip)
    assert r.validation.passed                          # coverage computed via re-parse
    assert r.validation.char_coverage >= CHAR_COVERAGE_THRESHOLD
    assert r.validation.output_table_cells == r.validation.source_table_cells


# --- registry wiring -------------------------------------------------------------
def test_document_formats_registered():
    for fmt in ("markdown", "txt", "html", "pdf", "docx", "pptx"):
        assert REGISTRY.get(fmt) is not None, fmt
        assert REGISTRY.get(fmt).modality == "document"
    assert REGISTRY.get("markdown").can_write          # md R+W
    assert not REGISTRY.get("pdf").can_write            # pdf read-only here (C19 writes)
