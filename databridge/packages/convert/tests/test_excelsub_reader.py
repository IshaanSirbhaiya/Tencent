"""C10 — Excel reader unit tests (09_FINAL_PRD §14). Per-requirement coverage of the
calamine+openpyxl reader against the Gauntlet Workbook fixture.

The EV-XL *gate* (cell-value 100%, header ≥95%, the swapped-epoch canary) lives in
``test_ev_xl.py``; this module pins each individual §14 behaviour."""
from __future__ import annotations

import datetime as _dt

import pytest

from databridge_convert.excelsub import (
    ExcelError,
    ExcelRefusal,
    ReadOptions,
    read_workbook,
)
from databridge_convert.excelsub.policies import ENCRYPTED_WORKBOOK
from databridge_convert.excelsub.provenance import cell_ref, col_letter, range_ref
from fixtures.excel.build import (
    MACRO_PAYLOAD,
    build_encrypted,
    build_gauntlet,
    build_macro_xlsm,
    EXPECTED,
)


@pytest.fixture(scope="module")
def gauntlet():
    return read_workbook(build_gauntlet(), options=ReadOptions(sheets="all"))


def _sheet(wb, name):
    return next(s for s in wb.sheets if s.name == name)


# --- §14.1 multiple sheets / selection ------------------------------------------
def test_all_sheets_present(gauntlet):
    assert [s.name for s in gauntlet.sheets] == EXPECTED["sheet_names_all"]


def test_visible_only_default_excludes_hidden_with_warning():
    wb = read_workbook(build_gauntlet())                      # default sheets="visible"
    assert [s.name for s in wb.sheets] == EXPECTED["sheet_names_visible"]
    assert any("hidden sheets excluded" in w for w in wb.warnings)


def test_explicit_sheet_selection():
    wb = read_workbook(build_gauntlet(), options=ReadOptions(sheets=["Sales"]))
    assert [s.name for s in wb.sheets] == ["Sales"]


def test_unknown_sheet_refused():
    with pytest.raises(ExcelRefusal) as ei:
        read_workbook(build_gauntlet(), options=ReadOptions(sheets=["Nope"]))
    assert ei.value.code == "SHEET_NOT_FOUND"


def test_hidden_sheet_flagged(gauntlet):
    assert _sheet(gauntlet, "_internal").visibility == "hidden"


# --- §14.3 typed cells + §14.5 cached formulas + §14.11 locale/leading-zeros -----
def test_typed_and_text_cells(gauntlet):
    row = _sheet(gauntlet, "Sales").cells[1]
    exp = EXPECTED["sales_row2"]
    assert row[0] == exp["id"]                                # "00012345" stays a STRING
    assert isinstance(row[0], str)
    assert row[2] == exp["amount_text"]                       # "1.234,56" stays text
    assert row[3] == exp["qty"] and isinstance(row[3], float)
    assert row[5] == exp["total"]                             # cached formula value
    assert row[6] == exp["created"] and isinstance(row[6], _dt.datetime)


def test_formula_flag_retained_with_cached_value(gauntlet):
    sales = _sheet(gauntlet, "Sales")
    assert "F2" in sales.formula_cells                        # has_formula
    assert sales.formula_cells["F2"].replace("=", "").upper() == "D2*E2"
    assert sales.cells[1][5] == 59.97                         # cached value used


def test_require_recalculated_refused():
    with pytest.raises(ExcelRefusal) as ei:
        read_workbook(build_gauntlet(), options=ReadOptions(require_recalculated=True))
    assert ei.value.code == "REQUIRE_RECALCULATED"


# --- §14.7 named ranges + §14.8 ListObjects -------------------------------------
def test_named_ranges(gauntlet):
    assert gauntlet.named_ranges_global["TaxRate"] == "Ledger!$H$2"
    assert "epoch_anchor" in gauntlet.named_ranges_global


def test_excel_table_listobject(gauntlet):
    assert "SalesTable" in _sheet(gauntlet, "Sales").tables


# --- §14.9 merged cells ----------------------------------------------------------
def test_data_merge_anchor_only(gauntlet):
    ledger = _sheet(gauntlet, "Ledger")
    assert "G2:G3" in ledger.merged_ranges
    # batch col (index 6); anchor G2 holds value, G3 stays None under anchor_only
    assert ledger.cells[1][6] == "BATCH-7"
    assert ledger.cells[2][6] is None


def test_header_merge_fill_down():
    # Quarterly's group row B1:C1="Q1", D1:E1="Q2" fills down for header context.
    wb = read_workbook(build_gauntlet(), options=ReadOptions(sheets=["Quarterly"]))
    q = wb.sheets[0]
    assert q.cells[0][1] == "Q1" and q.cells[0][2] == "Q1"    # filled across the merge
    assert q.cells[0][3] == "Q2" and q.cells[0][4] == "Q2"


def test_merged_refuse_policy():
    with pytest.raises(ExcelRefusal) as ei:
        read_workbook(build_gauntlet(),
                      options=ReadOptions(sheets=["Ledger"], merged_data="refuse"))
    assert ei.value.code == "MERGED_CELL_REFUSED"


# --- §14.10 hidden rows/cols -----------------------------------------------------
def test_hidden_column_flagged(gauntlet):
    # Ledger column K (index 10) is hidden.
    assert 10 in _sheet(gauntlet, "Ledger").hidden_cols


# --- §14.12 error cells ----------------------------------------------------------
def test_error_cells_null_warn_default(gauntlet):
    ledger = _sheet(gauntlet, "Ledger")
    assert ledger.error_cells == {"D2": "#N/A", "E2": "#VALUE!", "F2": "#DIV/0!"}
    # default policy null_warn → the grid cells are None
    assert ledger.cells[1][3] is None and ledger.cells[1][4] is None


def test_error_cells_keep_policy():
    wb = read_workbook(build_gauntlet(),
                       options=ReadOptions(sheets=["Ledger"], errors="keep"))
    cell = wb.sheets[0].cells[1][3]
    assert isinstance(cell, ExcelError) and cell.code == "#N/A"


def test_error_cells_hard_fail_policy():
    with pytest.raises(ExcelRefusal) as ei:
        read_workbook(build_gauntlet(),
                      options=ReadOptions(sheets=["Ledger"], errors="hard_fail"))
    assert ei.value.code == "ERROR_CELL"


# --- §14.6 ambiguous serials (NOT auto-converted) -------------------------------
def test_ambiguous_serial_column_flagged_not_converted(gauntlet):
    ledger = _sheet(gauntlet, "Ledger")
    assert 1 in ledger.ambiguous_date_cols                    # ship_date column
    assert ledger.cells[1][1] == 45000.0                      # stays a number, NOT a date


# --- §14.13 cross-sheet reference (cached) --------------------------------------
def test_cross_sheet_reference_cached(gauntlet):
    ledger = _sheet(gauntlet, "Ledger")
    assert ledger.cells[1][2] == EXPECTED["crosssheet_value"]
    assert "C2" in ledger.formula_cells


# --- §14.14 metadata -------------------------------------------------------------
def test_metadata_captured(gauntlet):
    assert gauntlet.metadata.created is not None
    assert gauntlet.metadata.has_macros is False


# --- §14.17 macros ---------------------------------------------------------------
def test_macros_detected_payload_never_leaks():
    wb = read_workbook(build_macro_xlsm())
    assert wb.metadata.has_macros is True
    assert any("macro" in x.lower() for x in wb.loss_report)
    dump = wb.model_dump_json().encode()
    assert MACRO_PAYLOAD not in dump
    assert b"FAKE_VBA" not in dump                            # no payload anywhere in IR


# --- §14.18 password-protected ---------------------------------------------------
def test_encrypted_workbook_refused():
    with pytest.raises(ExcelRefusal) as ei:
        read_workbook(build_encrypted())
    assert ei.value.code == ENCRYPTED_WORKBOOK
    assert ei.value.nearest_alternative                       # names a way forward


# --- §14.19 provenance helpers ---------------------------------------------------
def test_a1_provenance_helpers():
    assert col_letter(0) == "A" and col_letter(26) == "AA"
    assert cell_ref(6, 1) == "B7"
    assert range_ref(1, 0, 89, 3) == "A2:D90"


def test_workbook_cell_pointer_resolvable(gauntlet):
    from databridge_convert.excelsub import workbook_cell_pointer
    p = workbook_cell_pointer(gauntlet.sha256, "Sales", row0=1, col0=0)
    assert p.kind == "workbook_cell" and p.cell == "A2"
    assert p.workbook_sha == gauntlet.sha256


# --- coordinate-offset regressions (review H1/H3/H4a: sheets not starting at A1) -----
def _offset_workbook() -> bytes:
    """A sheet whose used range starts at C3 (NOT A1), with a hidden absolute column and a
    vertical data merge — exercises the absolute↔grid-local offset math."""
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Off"
    # header at C3:E3, data C4:E6
    for c, name in zip((3, 4, 5), ["sku", "secret", "qty"]):
        ws.cell(row=3, column=c, value=name)
    for i, (sku, sec, qty) in enumerate([("A", "x", 1), ("B", "y", 2), ("C", "z", 3)]):
        ws.cell(row=4 + i, column=3, value=sku)
        ws.cell(row=4 + i, column=4, value=sec)
        ws.cell(row=4 + i, column=5, value=qty)
    ws.column_dimensions["D"].hidden = True              # hide the absolute "secret" column
    # a vertical data merge spanning header+data in column G (must NOT fill into data, H1)
    ws["G3"] = "grp"
    ws["G3"].value = "GRP"
    ws.merge_cells("G3:G6")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_offset_sheet_origin_and_hidden_exclusion():
    from databridge_convert.excelsub import sheet_to_rowbatch
    wb = read_workbook(_offset_workbook(), options=ReadOptions(sheets="all"))
    s = wb.sheets[0]
    assert (s.row_origin, s.col_origin) == (2, 2)         # C3 → 0-based (2,2)
    assert 3 in s.hidden_cols                             # absolute index of column D
    rb = sheet_to_rowbatch(s, wb.sha256, include_hidden=False)
    # the hidden "secret" column (absolute D) is correctly excluded; "sku"/"qty" are kept
    assert "secret" not in rb.columns
    assert "sku" in rb.columns and "qty" in rb.columns
    rec0 = rb.records()[0]
    assert rec0["sku"] == "A" and rec0["qty"] == 1.0
    # provenance is ABSOLUTE A1 (sku lives in column C = abs col 2, data starts row 4)
    assert rb.column_provenance["sku"].cell_range.startswith("C4")


def test_vertical_data_merge_not_filled_into_data():
    # review H1: a tall merge reaching into data is a DATA merge (anchor_only) — its value
    # must NOT be filled down into data rows (that would corrupt values, §14 lossless).
    wb = read_workbook(_offset_workbook(), options=ReadOptions(sheets="all"))
    s = wb.sheets[0]
    # column G (abs 6) → grid-local 6-2=4; header row local 0 holds "GRP", data rows None
    assert s.cells[0][4] == "GRP"
    assert s.cells[1][4] is None and s.cells[2][4] is None
