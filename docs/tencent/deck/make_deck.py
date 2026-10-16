"""Build the pitch deck: HTML slides -> 1920x1080 PNGs -> PDF + PPTX (with speaker notes).

    python docs/tencent/deck/make_deck.py html      # write the slide HTML
    node   docs/tencent/deck/shoot_deck.cjs         # render PNGs
    python docs/tencent/deck/make_deck.py pack      # PDF + PPTX from the PNGs

The look and much of the story follow the Walnut customer deck (the company brain, how it works, the
business value, why Walnut wins); the hackathon slides add the fintech risk, how the Internal Brain answers
safely, the audit trail, the measured results and the Tencent tools. Outside numbers name their source on
the slide; our own numbers come from the repository; screens are real screenshots of the running app.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
OUT = REPO / "submission" / "07_deck"
SHOTS = REPO / "submission" / "03_app_scenes"
ASSETS = HERE / "assets"
RESULTS = REPO / "databridge" / "evals" / "internal_brain" / "results.json"

CSS = """
*{box-sizing:border-box}html,body{margin:0;width:1920px;height:1080px;overflow:hidden}
body{background:#F4F1EA;color:#16191F;font-family:Inter,'Segoe UI',system-ui,sans-serif}
.s{position:relative;height:100%;padding:76px 128px 110px;display:flex;flex-direction:column}
.k{color:#A8760C;letter-spacing:.22em;text-transform:uppercase;font-size:21px;font-weight:700}
h1{font-size:66px;line-height:1.06;margin:14px 0 14px;font-weight:800;letter-spacing:-.03em}
.sub{font-size:27px;color:#5E6470;line-height:1.45;margin:0 0 30px;max-width:1500px}
.card{background:#fff;border:1px solid #E6E1D6;border-radius:22px;padding:28px 30px;box-shadow:0 2px 0 rgba(0,0,0,.02)}
.navy{background:#212A45;color:#fff;border-color:#212A45}.navy .d{color:#C9CFDD}
.t{font-size:28px;font-weight:700;margin:0 0 10px;letter-spacing:-.01em}.d{font-size:21px;color:#5E6470;line-height:1.5}
.num{color:#A8760C;font-weight:700;font-size:18px;letter-spacing:.12em}
.big{font-size:118px;font-weight:800;letter-spacing:-.05em;line-height:1}.red{color:#E2583E}.ink{color:#212A45}
.src{font-size:16px;color:#8C919B;letter-spacing:.14em;text-transform:uppercase;margin-top:14px;font-weight:600}
.row{display:flex;gap:30px}.grid{display:grid;gap:26px}.g2{grid-template-columns:1fr 1fr}.g3{grid-template-columns:repeat(3,1fr)}
.g4{grid-template-columns:repeat(4,1fr)}.g5{grid-template-columns:repeat(5,1fr)}
.pill{display:inline-block;border:1px solid #DDD7CA;border-radius:999px;padding:12px 26px;margin:0 14px 0 0;font-size:22px;color:#3B414D;background:#FBFAF6;font-weight:500}
.shot{border:1px solid #E6E1D6;border-radius:18px;overflow:hidden;background:#0b1020;box-shadow:0 18px 40px rgba(33,42,69,.12)}
.shot img{width:100%;display:block}
.foot{position:absolute;left:128px;right:128px;bottom:40px;display:flex;justify-content:space-between;align-items:center;color:#8C919B;font-size:19px}
.brand{display:flex;align-items:center;gap:12px;color:#3B414D;font-weight:600;font-size:21px}.brand img{height:30px}
.hl{background:linear-gradient(transparent 62%,#F2DFA9 62%)}
.icon{width:46px;height:46px;border-radius:12px;background:#F6EAD0;display:flex;align-items:center;justify-content:center;color:#A8760C;font-size:22px;font-weight:700}
table{border-collapse:collapse;width:100%;font-size:21px}th{text-align:left;color:#5E6470;font-weight:600;padding:16px 14px;border-bottom:2px solid #DDD7CA}
td{padding:17px 14px;border-bottom:1px solid #E6E1D6;color:#3B414D}td.w{font-weight:700;color:#16191F}
td.us,th.us{background:#212A45;color:#fff;text-align:center;font-weight:700}
.tag{display:inline-block;font-size:16px;padding:6px 12px;border-radius:8px;background:#F6EAD0;color:#8A6110;font-weight:600;margin:3px 6px 3px 0}
"""


def page(n: int, total: int, body: str, foot: bool = True) -> str:
    footer = (f"<div class='foot'><span class='brand'><img src='{asset('walnut_icon.png')}'>Walnut Internal Brain</span>"
              f"<span>{n:02d} / {total}</span></div>") if foot else ""
    return ("<!doctype html><html lang='en'><head><meta charset='utf-8'><title>Slide</title>"
            "<link rel='preconnect' href='https://fonts.googleapis.com'>"
            "<link href='https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap' rel='stylesheet'>"
            f"<style>{CSS}</style></head><body><div class='s'>{body}{footer}</div></body></html>")


def asset(name: str) -> str:
    return (ASSETS / name).resolve().as_uri()


def shot(rel: str) -> str:
    return (SHOTS / rel).resolve().as_uri()


def head(kicker: str, title: str, sub: str = "") -> str:
    return (f"<div class='k'>{kicker}</div><h1>{title}</h1>" + (f"<p class='sub'>{sub}</p>" if sub else ""))


def card(title: str, text: str, num: str = "", cls: str = "card") -> str:
    n = f"<div class='num'>{num}</div>" if num else ""
    return f"<div class='{cls}'>{n}<div class='t'>{title}</div><div class='d'>{text}</div></div>"


def slides() -> list[tuple[str, str]]:
    h = json.loads(RESULTS.read_text(encoding="utf-8"))["headline"]
    n_req = f"{h['n_requests_all_modes']:,}"
    S: list[tuple[str, str]] = []

    S.append((f"""
<div style="display:flex;justify-content:space-between;align-items:center">
 <img src="{asset('walnut_wordmark.png')}" style="height:56px">
 <span style="color:#5E6470;font-size:22px;font-weight:500">Tencent Cloud AI CAN DO IT Hackathon · Singapore 2026</span></div>
<div style="margin-top:150px" class="k">FinTech Track · Aspire · The Internal Brain</div>
<div style="font-size:150px;font-weight:800;letter-spacing:-.05em;line-height:1;margin:24px 0 30px">Internal <span style="color:#D4911C">Brain</span></div>
<p class="sub" style="font-size:34px;max-width:1450px">Ask one question across Confluence, Jira, Slack and Google Drive. Get a cited answer built
<b>only from what you are allowed to open</b>, with every access on an audit trail nobody can quietly edit.</p>
<div><span class="pill">Permission-checked twice</span><span class="pill">Tamper-evident audit</span><span class="pill">Built with CodeBuddy · lives in WorkBuddy</span></div>
""", "Hi, we are Walnut. This is the Internal Brain, our entry for Aspire's FinTech challenge: one place to ask "
         "questions across the company's tools, answered only from what you are allowed to see, with every access "
         "on an audit trail.", False))

    S.append((f"""{head('01 · The reality', 'The data already exists')}
<div class="row" style="flex:1;align-items:center;gap:70px">
 <div style="flex:1"><div style="font-size:44px;line-height:1.35;font-weight:500"><span class="hl">The biggest obstacle to using advanced
 data analysis isn't skill base or technology;</span> it's plain old access to the data.</div><div class="src">Source · Harvard Business Review</div></div>
 <div style="flex:1" class="shot"><img src="{asset('source_hbr.png')}" style="background:#fff"></div></div>
""", "Every company already has the answers. The problem, as Harvard Business Review puts it, is plain old access."))

    S.append((f"""{head('02 · The cost', 'Searching takes a quarter of the week', 'The answers are somewhere in the organisation. Getting to them is where the time goes.')}
<div class="grid g3" style="margin-top:40px">
 <div><div class="big red">25%</div><div class="d" style="font-size:26px;margin-top:14px;color:#16191F">of a team's time goes to searching for answers.</div><div class="src">Atlassian · State of Teams 2025 · 12,000 knowledge workers</div></div>
 <div><div class="big ink">5.3<span style="font-size:52px"> hours</span></div><div class="d" style="font-size:26px;margin-top:14px;color:#16191F">a week, per person, waiting on colleagues for information or recreating work that already exists.</div><div class="src">Panopto with YouGov · 2018 · 1,001 employees</div></div>
 <div><div class="big ink">42%</div><div class="d" style="font-size:26px;margin-top:14px;color:#16191F">of an organisation's know-how is held only by individual employees. When they move on, it goes with them.</div><div class="src">Panopto with YouGov · 2018</div></div></div>
<div style="margin-top:auto;font-size:26px;font-weight:600;color:#212A45">The cost shows up as slower projects, repeated work and decisions made on partial information.</div>
""", "A quarter of the week goes on searching. Over five hours a week waiting on colleagues. And 42 percent of "
         "know-how sits in people's heads."))

    S.append((f"""{head('03 · The risk', 'AI projects stall before they start', 'Lack of AI-ready data puts AI projects at risk.')}
<div class="row" style="flex:1;align-items:center;gap:70px">
 <div style="flex:1"><div class="big red" style="font-size:190px">60%</div><div class="d" style="font-size:28px;color:#16191F;margin-top:10px">of AI projects unsupported by AI-ready data will be abandoned through 2026, Gartner forecasts.</div><div class="src">Source · Gartner newsroom · Feb 2025</div></div>
 <div style="flex:1" class="shot"><img src="{asset('source_gartner.png')}" style="background:#fff"></div></div>
""", "And AI does not fix it on its own: Gartner expects most AI projects without AI-ready data to be abandoned."))

    S.append((f"""{head('04 · The fintech edge', 'In a fintech, a leak can be a crime', 'An assistant that can read everything can also leak everything. For a payments company supervised by MAS, that blocks rollout entirely.')}
<div class="grid g4" style="margin-top:20px">
 {card('Tipping-off', 'Telling a customer that a suspicious-transaction report exists is an offence (CDSA s.57; MAS Notice PSN01).', '01')}
 {card('Card data', 'Card numbers pasted into a ticket must never reach an index or a model (PCI DSS 3.4.1).', '02')}
 {card('Oversharing and injection', 'Assistants oversharing, and instructions hidden in retrieved documents (EchoLeak, Slack AI), are on record.', '03')}
 {card('Stale permissions', 'Synced permissions lag the source by 30 minutes to a day. A removed person keeps seeing what they should not.', '04')}</div>
<div class="card navy" style="margin-top:34px"><div class="t" style="color:#fff">So the answer has to be useful and permission-safe at the same time.</div><div class="d">The model must never see anything the person asking cannot see, and compliance must be able to prove it afterwards.</div></div>
""", "In a fintech there is a second edge. A careless assistant can tip off a customer, leak card data or "
         "follow instructions hidden in a document. That is the problem Aspire set."))

    S.append((f"""{head('05 · What we heard', 'Teams told us where the hours go')}
<div class="grid g4" style="margin-top:16px">
 {card('Months before a project can start', 'Each new data or AI project spends six to eight months gathering data from scattered systems. The next project starts the hunt again.', '01')}
 {card('Access that arrives late', 'Data approvals drag on, so interns, new joiners and project teams sit waiting for the data they need.', '02')}
 {card('Tools chosen by old systems', 'Data held inside one vendor’s ecosystem keeps teams on that vendor’s tools, even when better ones exist.', '03')}
 {card('Knowledge that leaves with people', 'When a senior person moves on, the handover note holds only part of what they knew.', '04')}</div>
<div class="row" style="margin-top:34px;gap:26px">
 <div class="card" style="flex:1.2"><div class="num">ASPIRE · HACKATHON KICK-OFF</div><div style="font-size:30px;font-weight:600;margin-top:10px;line-height:1.35">"What can we offer this customer?" is the question that "takes time right now".</div><div class="src">Aspire's Head of AI</div></div>
 <div class="card" style="flex:1"><div class="num">OUR INTERVIEWS</div><div class="d" style="font-size:23px;margin-top:10px">Teams in aviation, asset management, IT services, insurance, sovereign wealth and health tech confirmed the same bottleneck.</div></div></div>
""", "We heard the same thing from teams across finance and beyond, and Aspire said it directly at the kick-off: "
         "answering 'what can we offer this customer' takes time right now."))

    S.append(("""<div style="margin:auto 0;text-align:center"><div class="k">What we do</div>
<div style="font-size:66px;font-weight:800;letter-spacing:-.03em;line-height:1.18;margin:30px auto 0;max-width:1500px">We make any organisation's <span style="color:#E2583E">scattered data</span> immediately usable: connected, understood and trusted, and only ever shown to the people <span style="color:#D4911C">allowed to see it.</span></div></div>
<div style="display:flex;justify-content:space-between;border-top:2px solid #E1DBCE;padding-top:22px;font-size:22px;color:#5E6470"><span>Category · <b style="color:#212A45">Enterprise knowledge, data and AI</b></span><span>The bottleneck is <b style="color:#212A45">safe access to data</b></span></div>
""", "That is what Walnut does: it makes scattered data usable, and in a fintech, safe to ask."))

    S.append(("""<div style="margin:auto 0"><div class="k">The solution</div>
<div style="font-size:120px;font-weight:800;letter-spacing:-.04em;line-height:1;margin:24px 0 34px">What is our<br>solution?</div>
<p class="sub" style="font-size:30px">Two layers on one foundation: a <b>Company Brain</b> that connects your data, and the <b>Internal Brain</b>
that answers each person only from what they are allowed to open, with every step on the record.</p></div>
""", "Two layers on one foundation: the company brain underneath, and the Internal Brain on top."))

    groups = [("Work tools", ["Confluence pages", "Jira issues", "Slack threads", "Google Drive files"]),
              ("Data platforms", ["Databases (SQL)", "Data-lake files (Parquet, Arrow, CSV)", "Warehouse exports", "Spreadsheets"]),
              ("Documents", ["PDFs, incl. scanned", "Word and PowerPoint", "Email and mailboxes", "Policies and runbooks"]),
              ("People", ["Meeting notes", "Handover notes", "Who can open what", "Contractors and guests"])]
    gh = "".join(f"<div class='card'><div class='num'>{g.upper()}</div>" + "".join(f"<div class='tag' style='font-size:19px'>{x}</div>" for x in xs) + "</div>" for g, xs in groups)
    S.append((f"""{head('Solution · Today', 'Company knowledge lives in a dozen places', 'Each team keeps its own records, in its own systems and formats, each with its own permission rules.')}
<div class="grid g4" style="margin-top:14px">{gh}</div>
<div class="card navy" style="margin-top:34px"><div class="t" style="color:#fff">Every system decides who may see what in its own way.</div><div class="d">Confluence page restrictions, Jira project roles and issue security, Slack private and Connect channels, Drive inheritance. An assistant has to respect all of them, live.</div></div>
""", "Today knowledge sits in a dozen places, each with its own permission model. That is what makes a safe assistant hard."))

    S.append((f"""{head('Solution · With Walnut', 'From scattered information to one record', 'Walnut connects information across your systems and brings it into one record, with every value pointing back to where it came from.')}
<div class="row" style="flex:1;align-items:center;gap:40px">
 <div style="flex:1" class="grid">
  {card('Confluence', 'DB Migration Plan (v6) · payment-gateway space')}
  {card('Jira', 'DBMIG-12 · Ledger cluster cutover (wave 4)')}
  {card('Slack', '#db-migration · "Replication lag on ledger-shard-3…"')}</div>
 <div style="width:150px;height:150px;border-radius:50%;background:#212A45;display:flex;flex-direction:column;align-items:center;justify-content:center;color:#fff;font-weight:700;font-size:20px;flex:none"><img src="{asset('walnut_icon.png')}" style="height:62px;filter:invert(1);margin-bottom:6px">Walnut</div>
 <div style="flex:1.2" class="card"><div class="num">ONE CONNECTED RECORD</div><div class="t" style="margin-top:8px">Database migration project</div>
  <div class="d"><div>Status · Wave 4: ledger cutover planned for 24 Oct <span class="tag">jira DBMIG-12</span></div>
  <div>Blocker · Replication lag on ledger-shard-3 <span class="tag">slack #db-migration</span></div>
  <div>Plan · Postgres 11 → 16, four waves <span class="tag">confluence cf-1001</span></div>
  <div style="margin-top:10px;color:#8C919B;font-size:18px">Every line opens to the source it came from, checked for the person asking.</div></div></div></div>
""", "Walnut pulls those pieces into one record. Here, a database migration: the plan from Confluence, the ticket "
         "from Jira, the blocker from Slack, each linked to its source."))

    steps = [("Connect", "Reads databases, documents, spreadsheets, email and the systems your teams use."),
             ("Understand", "Recognises the same customer, project or person across every source and links them into one record."),
             ("Convert", "Serves that record in the shape each job needs: answers, reports, analytics tables, AI context."),
             ("Prove", "Every figure points to the cell, line or page it came from. A value without proof is refused."),
             ("Protect", "Runs inside your environment with the model you choose. Each person sees only what they may open.")]
    st = "".join(f"<div class='card'><div class='num' style='text-align:right'>{i:02d}</div><div class='t'>{t}</div><div class='d'>{d}</div></div>"
                 for i, (t, d) in enumerate(steps, 1))
    S.append((f"""{head('Solution · How it works', 'How Walnut makes it usable and trusted')}
<div style="position:relative;height:40px;margin:30px 0 26px"><div style="position:absolute;top:19px;left:0;right:0;height:3px;background:#212A45"></div></div>
<div class="grid g5">{st}</div>
""", "Five steps: connect, understand, convert, prove and protect. The last two are what a fintech needs most."))

    S.append((f"""{head('Solution · Layer 1 · Company Brain', 'Inside the Company Brain', 'Walnut’s original screens, running on our synthetic fintech, Company A.')}
<div class="row" style="flex:1;gap:34px;min-height:0">
 <div style="flex:1.25" class="shot"><img src="{shot('walnut_original/03_brain_time_slider.png')}"></div>
 <div style="flex:1" class="grid">
  {card('Relationship map', 'How records connect: page to project, ticket to team, person to space.')}
  {card('Full history', 'When each fact was true and when it was recorded. Drag the time slider to rewind the company.')}
  {card('Source snapshots', 'A write-once copy of what each source said when read, so figures stay checkable.')}
  {card('Change tracking', 'When a source value changes, every answer built on it is flagged.')}</div></div>
""", "This is the company brain itself: a relationship map, full history you can rewind, source snapshots and "
         "change tracking. It is Walnut's original engine, reused unchanged."))

    flow = [("Who's asking", "A signed token. No prompt or tool argument can name a person."),
            ("Check 1 · index", "Search only what this person may open, with ranking statistics from those documents only."),
            ("Check 2 · live", "Ask Confluence, Jira, Slack or Drive again, now. Error or timeout means no."),
            ("Answer", "The model sees only passages that passed, and must cite one for every sentence."),
            ("Record", "Written to the sealed audit log before the answer is shown. Log down means no answer.")]
    fl = "".join(f"<div class='card'><div class='num'>{i:02d}</div><div class='t'>{t}</div><div class='d'>{d}</div></div>" for i, (t, d) in enumerate(flow, 1))
    S.append((f"""{head('Solution · Layer 2 · Internal Brain', 'How one question is answered safely', 'Two permission checks before the model sees anything, a grounding check after it answers, and an audit record before anything is shown.')}
<div class="grid g5" style="margin-top:10px">{fl}</div>
<div class="row" style="margin-top:30px;gap:26px">
 {card('Restricted looks exactly like non-existent', 'The refusal is the same bytes whether the thing is hidden from you or was never there, so nothing hints that more exists.', '', 'card navy')}
 {card('Your model, the same checks', 'Tencent Hy3 through TokenHub, or any OpenAI-compatible model. The checks do not depend on the model.')}</div>
""", "Every question goes through the same five steps. The key property: a restricted answer looks exactly like "
         "an answer about something that does not exist."))

    S.append((f"""{head('In practice', 'Same question, two people', 'Jordan, a backend engineer, and Ken, an outside contractor, ask the same thing.')}
<div class="row" style="flex:1;gap:34px;min-height:0">
 <div style="flex:1.6" class="shot"><img src="{shot('internal_brain/12_compare.png')}"></div>
 <div style="flex:1" class="grid">
  {card('Jordan gets the full picture', 'A cited answer from Slack, Jira and Confluence, every source re-checked live.', '1')}
  {card('Ken gets a polite no', '"I couldn’t find an answer to that in the sources you can access." No hint more exists.', '2')}
  {card('Both are logged', 'Each question, decision and answer is on the audit trail.', '3')}</div></div>
""", "This is the real app. Same question, two people. Jordan gets the answer; Ken gets a polite no, with no hint "
         "that anything is hidden."))

    S.append((f"""{head('In practice · Compliance', 'Who asked what, and proof nobody edited it')}
<div class="row" style="flex:1;gap:34px;min-height:0">
 <div style="flex:1.6" class="shot"><img src="{shot('internal_brain/17_audit_verify.png')}"></div>
 <div style="flex:1" class="grid">
  {card('Ask the audit log in plain English', '"Everything Jordan accessed in the payment-gateway space, last 30 days" becomes a visible filter.')}
  {card('A reason to read an answer', 'Compliance sees metadata by default; reading an answer needs a reason, and that is logged too.')}
  <div class="card navy"><div style="font-size:64px;font-weight:800;color:#F2C46B">6 / 6</div><div class="d">tamper attacks caught by "Verify audit chain": edits, deletions, rollback, a key-holder rewrite.</div></div></div></div>
""", "For compliance: plain-English audit search, a reasoned reveal, and a sealed log. Our tamper demo edits, "
         "deletes and rolls back the log, and verification catches all six attacks."))

    S.append((f"""{head('Proof', 'Checked against answers we already knew', 'We built a synthetic fintech, Company A, with 1,143 documents and 49 planted facts, so every right answer was known in advance.')}
<div class="grid g4" style="margin-top:30px">
 <div><div class="big red" style="font-size:104px">0</div><div class="t" style="margin-top:10px">leaks in {n_req} attacks</div><div class="d">Restricted content in any answer, prompt or response.</div><div class="src">Leak evaluation · Oct 2026</div></div>
 <div><div class="big ink" style="font-size:104px">0</div><div class="t" style="margin-top:10px">differences in 400 pairs</div><div class="d">"Restricted" and "doesn't exist" give the same answer.</div><div class="src">Differential test</div></div>
 <div><div class="big ink" style="font-size:104px">0</div><div class="t" style="margin-top:10px">injection failures</div><div class="d">Instructions planted on all four platforms were treated as content.</div><div class="src">976 injection requests</div></div>
 <div><div class="big ink" style="font-size:104px">677</div><div class="t" style="margin-top:10px">automated tests</div><div class="d">639 on the Internal Brain, 38 in the browser, all passing.</div><div class="src">Repository test suites</div></div></div>
<div style="margin-top:auto;font-size:25px;font-weight:600;color:#212A45">The harness ships with the code, so a pilot runs the same checks on your own permissions.</div>
""", "We attacked it on purpose: almost six thousand adversarial requests with zero leaks, four hundred paired "
         "answers with zero differences, and every scenario covered by tests."))

    uc = [("“What can we offer this customer?” takes time.", "Approved, market-specific answers for ops and account managers, with sources.", "Faster customer answers"),
          ("Account managers must not tip off a customer.", "They get the approved script, never the case notes, worded like any other answer.", "No tipping-off risk"),
          ("Access reviews take weeks of log digging.", "Compliance asks the audit log in plain English, with tamper-proof records.", "Audits in minutes"),
          ("Contractors and new joiners wait for access.", "Access follows each person's permissions at the source, checked live.", "Safe, faster onboarding"),
          ("Know-how leaves when people move on.", "Decisions and context stay in the record with the evidence behind them.", "Knowledge that stays")]
    rows = "".join(f"<tr><td class='w'>{a}</td><td>{b}</td><td class='us'>{c}</td></tr>" for a, b, c in uc)
    S.append((f"""{head('Use cases', 'Where a fintech gets its time back')}
<table style="margin-top:10px"><tr><th style="width:36%">What slows teams down today</th><th>What changes with the Internal Brain</th><th class="us" style="width:22%">What you get back</th></tr>{rows}</table>
""", "Where Aspire gets its time back: faster customer answers, no tipping-off risk, audits in minutes, safe onboarding."))

    S.append((f"""{head('Security and control', 'Your data stays inside your walls')}
<div class="grid g3" style="margin-top:12px">
 {card('Runs in your environment', 'On your machines or your cloud. Embeddings are local, so nothing leaves at ingest.')}
 {card('Your choice of model', 'Tencent Hy3 via TokenHub, any OpenAI-compatible model, or none: the quoted-answer path needs no model.')}
 {card('Access checked twice, live', 'Each person sees only what the source system lets them open, re-checked at answer time.')}
 {card('A person stays in control', 'Compliance must give a reason to read an answer, and that reveal is logged.')}
 {card('Sealed, verifiable log', 'Chained and sealed, with its head published off-host. Edits and deletions are detected.')}
 {card('Sensitive values masked', 'Card numbers, CVVs, IBANs and NRICs are masked before indexing, so no model ever sees them.')}</div>
<div style="margin-top:auto;font-size:24px;font-weight:600;color:#212A45">Mapped to MAS TRM, PCI DSS and CDSA, with the limits stated: docs/tencent/MAS_TRM_MAPPING.md.</div>
""", "Security and control: it runs in your environment, with your choice of model, checks access twice, keeps a "
         "person in control and masks sensitive values."))

    tools = [("CodeBuddy", "Six build sessions: trust boundaries, the WorkBuddy connectors, the token tool, model-chain tests, a security review, then fixes to its own code."),
             ("WorkBuddy", "The second front door: connectors that ask as Jordan, as Ken, and for compliance, each with its own signed token."),
             ("Hunyuan Hy3 · TokenHub", "First in the answer chain, with DeepSeek as backup. Wired in and tested; live calls wait on account verification."),
             ("Lighthouse", "A hardened one-command deploy kit: HTTPS, passcode-gated demo, only ports 80 and 443 open."),
             ("Miora", "Made the cover art from our written brief.")]
    tl = "".join(card(t, d, f"{i:02d}") for i, (t, d) in enumerate(tools, 1))
    S.append((f"""{head('Built with Tencent', 'Tencent tools, used for real', 'Every claim here has a transcript, screenshot or test in the repository.')}
<div class="grid g5" style="margin-top:10px">{tl}</div>
<div class="card navy" style="margin-top:34px"><div class="t" style="color:#fff">93% of the code is the original Walnut engine, reused unchanged.</div><div class="d">The Internal Brain, its two checks, the audit console, the web app, the fintech rules and the WorkBuddy connectors were built during the hackathon.</div></div>
""", "Built with Tencent tools throughout: CodeBuddy for the build, WorkBuddy as a front door, Hy3 through "
         "TokenHub for answers, a Lighthouse deploy kit and Miora for the cover."))

    S.append((f"""{head('Business value', 'Measure it in hours saved and risk removed')}
<div class="row" style="gap:30px;flex:1;min-height:0">
 <div class="card" style="flex:1"><div class="num">TIME BACK</div>
  <div class="row" style="gap:14px;margin:20px 0;align-items:center;font-size:21px;font-weight:600"><div class="card" style="padding:16px 20px;background:#F7F4EC">Hours per person per week</div>×<div class="card" style="padding:16px 20px;background:#F7F4EC">People</div>×<div class="card" style="padding:16px 20px;background:#F7F4EC">Loaded hourly cost</div></div>
  <div style="background:#212A45;color:#fff;border-radius:14px;padding:16px 22px;font-size:22px;font-weight:600">= Weekly cost of searching, waiting and rework</div>
  <div class="src" style="margin-top:26px">Example · a 600-person fintech</div>
  <div class="d" style="font-size:23px;color:#16191F;margin-top:8px">Saving 1.5 hours a week at US$31 an hour is about <b class="red">US$1.3M a year</b>. Break-even for a US$100k deployment is <b>7 minutes a week</b> per person.</div>
  <div class="d" style="font-size:18px;margin-top:10px">A projection, not a measurement. We measure it with you on one workflow, against your own baseline.</div></div>
 <div class="card navy" style="flex:1"><div class="num" style="color:#F2C46B">RISK REMOVED</div>
  <div class="t" style="color:#fff;margin-top:18px">An assistant MAS supervision can accept</div><div class="d">Without permission-aware retrieval and an audit trail, an internal AI assistant cannot be rolled out at all.</div>
  <div class="t" style="color:#fff;margin-top:22px">Prepared customer conversations</div><div class="d">Account history, approved offers and open issues, cited, before every call.</div>
  <div class="t" style="color:#fff;margin-top:22px">Audits that take minutes</div><div class="d">Every access reconstructable, with proof the record was not altered.</div></div></div>
""", "The value is time back, about 1.3 million US dollars a year for a 600-person fintech in our base case, and "
         "risk removed: without this, an internal assistant cannot be rolled out under MAS supervision."))

    cmp_rows = [("Answers only from what each person may open, checked live", "no", "synced ACLs", "no", "synced ACLs"),
                ("“Restricted” indistinguishable from “doesn't exist”", "no", "no", "no", "no"),
                ("Tamper-evident audit of every question and answer", "partial", "no", "no", "logs"),
                ("Connects databases, documents, chat and tickets", "partial", "partial", "no", "✓"),
                ("Every figure traces to its cell, line or page", "lineage only", "chunks", "features", "platform"),
                ("Runs in your environment with any model", "cloud-bound", "varies", "✓", "enterprise")]
    tr = "".join(f"<tr><td class='w'>{r[0]}</td>" + "".join(f"<td style='text-align:center'>{c}</td>" for c in r[1:]) + "<td class='us'>✓</td></tr>" for r in cmp_rows)
    S.append((f"""{head('Differentiation', 'Why Walnut wins')}
<table style="margin-top:6px"><tr><th style="width:38%">Capability</th><th style="text-align:center">Data warehouses</th><th style="text-align:center">Search / RAG tools</th><th style="text-align:center">Feature stores</th><th style="text-align:center">Large AI platforms</th><th class="us">Walnut</th></tr>{tr}</table>
<div style="margin-top:auto;font-size:25px;font-weight:600;color:#212A45">Only Walnut spans the Company Brain, live permission checks and an audit trail you can prove.</div>
""", "Why Walnut wins: live permission checks, restricted indistinguishable from non-existent, a tamper-evident "
         "audit, and evidence behind every figure, in one system."))

    team = [("team_dev.png", "Dev", "CS + Business @ NTU CCDS", "AI and Automation at Singlife · AI/ML Engineer at Neurowyzr · AI and Automation at Bajaj Finserv · Researcher, NTU-Schaeffler Lab · Tech Lead, NTU AI&ML Club."),
            ("team_mehul.png", "Mehul", "CS @ NTU", "Production AI agent over 16,000+ records at Changi's AI/ML Lab · Privacy-aware clinical AI at Neurowyzr · Multi-agent AI research at A*STAR."),
            ("team_aaruni.png", "Aaruni", "USTEP & ASD @ SUTD", "Research Intern at SEC · Architecture Intern at HighSpan, MAK Designs, TAS · Design Intern at WGS."),
            ("team_ishaan.png", "Ishaan", "CSD @ SUTD", "Real-time power-grid and industrial cybersecurity research at SUTD's iTrust Centre · Production full-stack systems at Bajaj Finance and client projects.")]
    tm = "".join(f"""<div style="text-align:center"><div style="width:190px;height:190px;border-radius:50%;overflow:hidden;margin:0 auto 22px;border:6px solid #fff;box-shadow:0 10px 30px rgba(0,0,0,.12)"><img src="{asset(f)}" style="width:100%;height:100%;object-fit:cover"></div>
<div class="t" style="font-size:34px">{n}</div><div class="num">{r.upper()}</div><div class="d" style="margin-top:12px">{b}</div></div>""" for f, n, r, b in team)
    S.append((f"""<div style="display:flex;justify-content:space-between"><div class="k">The team</div><div class="k">Backed by BSF · SUTD VIE</div></div>
<h1>Built by people who live the data problem</h1>
<div class="grid g4" style="margin-top:40px;gap:44px">{tm}</div>
""", "We are Dev, Mehul, Aaruni and Ishaan, and we have lived this problem in finance, insurance, design and security work."))

    nx = [("See it on your permissions", "Run our leak harness against Aspire's own permission model in a sandbox."),
          ("Agree the workflow", "Pick one: the “what can we offer this customer” question, or compliance access reviews."),
          ("Connect in your environment", "Confluence, Jira, Slack and Drive, with Hy3 through TokenHub."),
          ("Measure", "Hours saved and audit time against your own baseline."),
          ("Decide what comes next", "Expand to more teams, with the full record always yours to keep.")]
    nxt = "".join(f"<div><div style='width:22px;height:22px;border-radius:50%;background:{'#D4911C' if i == 5 else '#212A45'};margin:-11px 0 26px'></div><div class='num'>{i:02d}</div><div class='t' style='font-size:25px{';color:#A8760C' if i == 5 else ''}'>{t}</div><div class='d'>{d}</div></div>" for i, (t, d) in enumerate(nx, 1))
    S.append((f"""{head('Next step', 'Start with one workflow Aspire repeats', 'Pick one job that costs your team hours every week. We set it up inside your environment, measure the result against your baseline, and you decide what comes next.')}
<div style="border-top:3px solid #212A45;margin-top:60px;padding-top:0" class="grid g5">{nxt}</div>
<div style="margin-top:auto;display:flex;justify-content:space-between;border-top:2px solid #E1DBCE;padding-top:22px;font-size:24px"><b>github.com/IshaanSirbhaiya/Tencent</b><span style="color:#5E6470">Start at INTERNAL_BRAIN.md · demo video in the submission</span></div>
""", "Next step: a pilot on one workflow, starting with our leak harness run against Aspire's own permissions. Thank you."))
    return S


def write_html() -> None:
    S = slides()
    (OUT / "html").mkdir(parents=True, exist_ok=True)
    for old in (OUT / "html").glob("slide*.html"):
        old.unlink()
    notes = {}
    total = len(S)
    for i, item in enumerate(S, 1):
        body, note = item[0], item[1]
        foot = item[2] if len(item) > 2 else True
        (OUT / "html" / f"slide{i:02d}.html").write_text(page(i, total, body, foot), encoding="utf-8")
        notes[f"slide{i:02d}"] = note
    (OUT / "notes.json").write_text(json.dumps(notes, indent=1, ensure_ascii=False), encoding="utf-8")
    print(f"{total} slides -> {OUT / 'html'}")


def pack() -> None:
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Emu
    notes = json.loads((OUT / "notes.json").read_text(encoding="utf-8"))
    pngs = sorted((OUT / "png").glob("slide*.png"))
    imgs = [Image.open(p).convert("RGB") for p in pngs]
    pdf = OUT / "walnut_internal_brain_deck.pdf"
    imgs[0].save(pdf, save_all=True, append_images=imgs[1:], resolution=144)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)          # 16:9, 13.333 x 7.5 in
    blank = prs.slide_layouts[6]
    for p in pngs:
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(str(p), 0, 0, width=prs.slide_width, height=prs.slide_height)
        s.notes_slide.notes_text_frame.text = notes.get(p.stem, "")
    pptx_path = OUT / "walnut_internal_brain_deck.pptx"
    prs.save(pptx_path)
    print("->", pdf, "\n->", pptx_path)


if __name__ == "__main__":
    {"html": write_html, "pack": pack}[sys.argv[1] if len(sys.argv) > 1 else "html"]()
