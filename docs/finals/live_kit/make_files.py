"""Build the live-upload demo files for the Walnut DataBridge live-ingestion kit.

Synthetic only: Company A (companya.example), its people and its Jira keys come from
databridge/fixtures/enterprise/world.json. No real company, person or card appears.

    databridge/.venv/Scripts/python.exe docs/finals/live_kit/make_files.py

Writes six files into docs/finals/live_kit/files/ and copies them to
databridge/apps/internal_brain_web/public/finals/live_kit/files/ so the browser can fetch them.
Formats: .docx (python-docx), .xlsx (openpyxl), .pptx (python-pptx), .txt, .csv, .md. No PDF
(scope: only formats the existing readers already handle with what is installed).
Each file carries ONE clear key fact (see live_samples.json, field ``proof_phrase``).
"""
from __future__ import annotations

import csv
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.util import Emu, Inches, Pt as PPt

HERE = Path(__file__).resolve().parent
OUT = HERE / "files"
REPO = HERE.parents[2]
PUBLIC = REPO / "databridge" / "apps" / "internal_brain_web" / "public" / "finals" / "live_kit" / "files"

INK = "1F2A37"
ACCENT = "2F5D8A"
MUTED = "6B7785"
FOOTER = "Company A internal - synthetic demonstration data, invented for Walnut DataBridge."

DOCX = "DB Migration — Wave 4 cutover plan v2.docx"
XLSX = "Q4 FX fee schedule — SG & ID (draft for review).xlsx"
PPTX = "Payment outage — post-incident review.pptx"
CHAT = "WhatsApp Chat with Ops on-call.txt"
CSVF = "vendor-sla-tracker.csv"
MD = "Onboarding buddy guide.md"
ALL = (DOCX, XLSX, PPTX, CHAT, CSVF, MD)


# ------------------------------------------------------------------------------------- Word
def make_docx(path: Path) -> None:
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    title = d.add_heading("DB Migration — Wave 4 cutover plan v2", level=0)
    for r in title.runs:
        r.font.color.rgb = RGBColor.from_string(INK)
    meta = d.add_paragraph()
    meta.add_run("Owner: ").bold = True
    meta.add_run("Priya Raman (priya.sre), cutover operations    ")
    meta.add_run("Status: ").bold = True
    meta.add_run("Approved for rehearsal    ")
    meta.add_run("Ticket: ").bold = True
    meta.add_run("DBMIG-12")
    for r in meta.runs:
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor.from_string(MUTED)

    d.add_heading("1. What changed since v1", level=1)
    d.add_paragraph(
        "The wave 4 ledger cutover (DBMIG-12) is rescheduled to Saturday 14 November 2026, "
        "22:00 SGT. Version 1 of this plan said 24 October 2026.")
    d.add_paragraph(
        "Why: the ledger-shard-3 copy job was paused at batch 1,812 after replication lag passed "
        "40 seconds, and the refund idempotency fix (DBMIG-14) has to land before the final "
        "backfill. The overall approach is unchanged; see the DB Migration Plan: "
        "https://companya.atlassian.net/wiki/spaces/payment-gateway/pages/cf-1001")

    d.add_heading("2. Cutover window", level=1)
    for line in ("Start: Saturday 14 November 2026, 22:00 SGT (payments traffic lowest).",
                 "Write freeze: 22:00 to 23:30 SGT; card authorisations keep flowing through the "
                 "dual-write shim.",
                 "Go / no-go call: 21:30 SGT, chaired by Nina Kowalski.",
                 "Fallback date if the call says no-go: Saturday 21 November 2026, same times."):
        d.add_paragraph(line, style="List Bullet")

    d.add_heading("3. Steps", level=1)
    rows = [
        ("Step", "Time (SGT)", "Action", "Owner"),
        ("1", "21:30", "Go / no-go call; confirm replica lag under 5 s on all four ledger shards", "nina.k"),
        ("2", "22:00", "Freeze ledger writes on the Postgres 11 cluster; dual-write shim to drain mode", "jdoe"),
        ("3", "22:10", "Final backfill delta for ledger shards 1-4; verify row counts and checksums", "priya.sre"),
        ("4", "22:45", "Flip the ledger read path to the managed Postgres 16 cluster", "priya.sre"),
        ("5", "23:00", "Smoke tests: authorisation, refund, settlement file dry run", "wei.l"),
        ("6", "23:30", "Lift the write freeze; old cluster stays read-only for 7 days", "priya.sre"),
    ]
    t = d.add_table(rows=len(rows), cols=4)
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.text = v
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    r.bold = (i == 0)

    d.add_heading("4. Rollback", level=1)
    d.add_paragraph(
        "Until step 6, the Postgres 11 cluster remains the source of truth. Rollback is a flag "
        "flip of the dual-write shim back to primary-old, then a replay of the delta captured "
        "after step 2. Decision owner: Priya Raman.")

    d.add_heading("5. Sign-off", level=1)
    d.add_paragraph("Nina Kowalski (plan owner), Priya Raman (cutover operations), "
                    "Jordan Doe (dual-write shim), Yuki Sato (data platform).")
    foot = d.add_paragraph(FOOTER)
    for r in foot.runs:
        r.font.size = Pt(8)
        r.italic = True
        r.font.color.rgb = RGBColor.from_string(MUTED)
    d.core_properties.author = "Priya Raman"
    d.core_properties.title = "DB Migration — Wave 4 cutover plan v2"
    d.save(path)


# ------------------------------------------------------------------------------------ Excel
def make_xlsx(path: Path) -> None:
    wb = Workbook()
    thin = Side(style="thin", color="C9D1DA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor=ACCENT)
    new_fill = PatternFill("solid", fgColor="FFF4CC")

    def sheet(ws, market: str, rows: list[tuple], note: str, extra: list[str]) -> None:
        ws.title = market
        ws["A1"] = f"Q4 FX fee schedule — {market} SME plans (DRAFT for review, not approved)"
        ws["A1"].font = Font(bold=True, size=13, color=INK)
        ws["A2"] = ("Prepared by Leo Hartono (leo.ops), Product & Operations. "
                    "Reviewers: Hana Yusuf, Benjamin Chua.")
        ws["A2"].font = Font(italic=True, size=9, color=MUTED)
        hdr = ("Plan", "Product", "Current fee (v3)", "Proposed Q4 fee", "Effective from", "Notes")
        for j, h in enumerate(hdr, 1):
            c = ws.cell(row=4, column=j, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for i, r in enumerate(rows, 5):
            for j, v in enumerate(r, 1):
                c = ws.cell(row=i, column=j, value=v)
                c.border = border
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if j == 4:
                    c.fill = new_fill
                    c.font = Font(bold=True)
        n = 5 + len(rows) + 1
        ws.cell(row=n, column=1, value=note).font = Font(bold=True, color=INK)
        for k, line in enumerate(extra, 1):
            ws.cell(row=n + k, column=1, value=line).font = Font(size=9, color=MUTED)
        ws.cell(row=n + len(extra) + 2, column=1, value=FOOTER).font = Font(italic=True, size=8,
                                                                           color=MUTED)
        for j, wdt in enumerate((14, 26, 18, 18, 16, 46), 1):
            ws.column_dimensions[get_column_letter(j)].width = wdt
        ws.freeze_panes = "A5"

    sg = wb.active
    sheet(sg, "SG", [
        ("Growth", "FX conversion", "0.35%", "0.30%", "1 Dec 2026",
         "Draft: Growth SME plan only; no minimum charge"),
        ("Starter", "FX conversion", "0.35%", "0.35%", "-", "No change proposed"),
        ("Growth", "SWIFT transfer", "SGD 10", "SGD 8", "1 Dec 2026", "Matches the ID corridor review"),
        ("All", "Local SGD transfer", "Free", "Free", "-", "No change"),
    ], "Key proposal: the Q4 draft cuts the SG FX conversion fee for the Growth SME plan from "
       "0.35% to 0.30%, effective 1 December 2026, if approved.",
        ["Masking test (a published test card number, not a real card): 4000 0566 5566 5556",
         "Approval: pricing committee on 18 November 2026. Until then fee schedule v3 applies."])
    idw = wb.create_sheet()
    sheet(idw, "ID", [
        ("Growth", "Multi-currency account FX margin", "0.40%", "0.38%", "1 Dec 2026",
         "Draft: USD, SGD and IDR wallets"),
        ("Starter", "Multi-currency account", "not offered", "not offered", "-", "No change"),
        ("Growth", "SWIFT transfer", "USD 12", "USD 10", "1 Dec 2026", "Draft"),
    ], "Key proposal: the ID Growth plan FX margin moves from 0.40% to 0.38% in the Q4 draft.",
        ["Approval: pricing committee on 18 November 2026."])
    wb.properties.creator = "Leo Hartono"
    wb.properties.title = "Q4 FX fee schedule — SG & ID (draft for review)"
    wb.save(path)


# ------------------------------------------------------------------------------- PowerPoint
def make_pptx(path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def band(slide, title: str, kicker: str) -> None:
        bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = PRGB.from_string(ACCENT)
        bar.line.fill.background()
        tb = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.5)).text_frame
        tb.text = kicker.upper()
        tb.paragraphs[0].runs[0].font.size = PPt(12)
        tb.paragraphs[0].runs[0].font.color.rgb = PRGB.from_string(MUTED)
        tt = slide.shapes.add_textbox(Inches(0.7), Inches(0.85), Inches(12), Inches(1)).text_frame
        tt.text = title
        r = tt.paragraphs[0].runs[0]
        r.font.size, r.font.bold = PPt(30), True
        r.font.color.rgb = PRGB.from_string(INK)
        ft = slide.shapes.add_textbox(Inches(0.7), Inches(6.95), Inches(12), Inches(0.4)).text_frame
        ft.text = FOOTER
        ft.paragraphs[0].runs[0].font.size = PPt(9)
        ft.paragraphs[0].runs[0].font.color.rgb = PRGB.from_string(MUTED)

    def bullets(slide, lines: list[str], top=2.0, size=18) -> None:
        tf = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11.6), Inches(4.6)).text_frame
        tf.word_wrap = True
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = "•  " + line
            p.space_after = PPt(10)
            p.runs[0].font.size = PPt(size)
            p.runs[0].font.color.rgb = PRGB.from_string(INK)

    def table(slide, rows: list[tuple], widths: list[float], top=2.0) -> None:
        shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.7), Inches(top),
                                     Inches(sum(widths)), Inches(0.5 * len(rows)))
        tbl = shp.table
        for j, w in enumerate(widths):
            tbl.columns[j].width = Emu(int(Inches(w)))
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                cell = tbl.cell(i, j)
                cell.text = v
                run = cell.text_frame.paragraphs[0].runs[0]
                run.font.size = PPt(14)
                run.font.bold = (i == 0)

    s = prs.slides.add_slide(blank)
    band(s, "Payment outage — post-incident review", "Payments engineering · held 28 October 2026")
    bullets(s, ["Incident: card authorisations failed intermittently for 47 minutes on 19 August 2026.",
                "Review chaired by Rajesh Menon; notes by Jordan Doe; attendees from Payments, SRE "
                "and Support.",
                "RCA page: https://companya.atlassian.net/wiki/spaces/payment-gateway/pages/cf-1002"],
            top=2.3)

    s = prs.slides.add_slide(blank)
    band(s, "Timeline (SGT, 19 August 2026)", "What happened")
    table(s, [("Time", "Event"),
              ("14:02", "First soft declines at the card-network gateway"),
              ("14:09", "Support escalates merchant complaints"),
              ("14:21", "SRE paged (support noticed before monitoring)"),
              ("14:38", "Expired intermediate TLS certificate identified"),
              ("14:49", "New certificate deployed; authorisations recover")], [1.6, 10.2])

    s = prs.slides.add_slide(blank)
    band(s, "Root cause", "Why it happened")
    bullets(s, ["Certificate renewal for the card-network gateway was a manual calendar task.",
                "The reminder went to a deactivated mailbox of an engineer who had left.",
                "No expiry monitoring on outbound mTLS endpoints; the failover path shared the "
                "same chain."])

    s = prs.slides.add_slide(blank)
    band(s, "Follow-up tickets", "What we are doing about it")
    table(s, [("Ticket", "Action", "Owner", "Status"),
              ("PAY-201", "Certificate expiry alerts for every outbound mTLS endpoint", "tom.b",
               "In Progress"),
              ("PAY-202", "Move card-network certificate renewal to the cert manager", "tom.b",
               "In Progress"),
              ("PAY-203", "Separate certificate chain for the failover path", "tom.b", "In Progress"),
              ("PAY-204", "Alert on authorisation success rate below 95% for 5 minutes", "tom.b",
               "Done"),
              ("PAY-205", "Audit ownership of shared mailboxes and calendar reminders", "tom.b",
               "In Progress")],
          [1.5, 7.2, 1.4, 1.8])

    s = prs.slides.add_slide(blank)
    band(s, "Decision from the post-incident review", "Deadline")
    bullets(s, ["Decision: all five follow-up tickets, PAY-201 to PAY-205, must close by "
                "30 November 2026.",
                "PAY-203 is the riskiest: failover chain separation needs a maintenance window.",
                "Next review: 4 December 2026, same attendees."], size=22)

    s = prs.slides.add_slide(blank)
    band(s, "What went well, and what we keep", "Lessons")
    bullets(s, ["Rollback tooling worked; the status page was updated within 15 minutes.",
                "PAY-204 is live: the success-rate alert now pages before support has to.",
                "Owners named for every renewal; no reminder goes to a personal mailbox."])
    prs.core_properties.author = "Jordan Doe"
    prs.core_properties.title = "Payment outage — post-incident review"
    prs.save(path)


# --------------------------------------------------------------------------- WhatsApp (.txt)
# Android export format, one message per line; the key fact is the 09:15 line from Priya.
CHAT_LINES = [
    ("24/10/2026, 02:14", "Devansh Patel", "Paged: authorisation success rate under 95% for 5 min. Looking now"),
    ("24/10/2026, 02:15", "Devansh Patel", "Dashboard shows a spike of retries from the settlement side, not the card network"),
    ("24/10/2026, 02:17", "Aisha Rahman", "I'm up. Card-network gateway certificates all fine, checked expiry"),
    ("24/10/2026, 02:19", "Devansh Patel", "Status page? Merchants not complaining yet"),
    ("24/10/2026, 02:20", "Aisha Rahman", "Holding off, success rate already back at 98%"),
    ("24/10/2026, 02:26", "Devansh Patel", "Recovered fully at 02:24. Will write it up in the morning"),
    ("24/10/2026, 08:41", "Priya Raman", "Morning. Read the thread. Was this the same alert from PAY-204?"),
    ("24/10/2026, 08:43", "Devansh Patel", "Yes, the PAY-204 success-rate alert. First real page since it shipped"),
    ("24/10/2026, 08:50", "Jordan Doe", "Settlement job logs show a burst at 02:10. Checking the retry config"),
    ("24/10/2026, 09:12", "Jordan Doe", "Found it: retry backoff was set to 0 after Thursday's config change"),
    ("24/10/2026, 09:15", "Priya Raman", "Root cause for Saturday's alert: the settlement retry job resent 3,400 "
                                         "authorisations in one burst at 02:10, which tripped the PAY-204 "
                                         "success-rate alert; no customer impact."),
    ("24/10/2026, 09:16", "Priya Raman", "Jordan can you restore the backoff today?"),
    ("24/10/2026, 09:18", "Jordan Doe", "On it. PR up in 30 min, backoff back to 2s with jitter"),
    ("24/10/2026, 09:40", "Aisha Rahman", "Should we raise the alert threshold so retries don't page?"),
    ("24/10/2026, 09:42", "Priya Raman", "No. It did its job. Keep PAY-204 as is"),
    ("24/10/2026, 10:05", "Jordan Doe", "Backoff fix merged and deployed to prod-sg"),
    ("24/10/2026, 10:06", "Devansh Patel", "Thanks. I'll watch tonight's settlement run"),
    ("25/10/2026, 02:20", "Devansh Patel", "Settlement run clean. No retries above baseline"),
    ("25/10/2026, 02:21", "Aisha Rahman", "Nice. Going back to sleep"),
    ("25/10/2026, 09:30", "Priya Raman", "Writing the note for Monday standup. Anything else to add?"),
    ("25/10/2026, 09:34", "Jordan Doe", "Add a config check so backoff can't be 0 again"),
    ("25/10/2026, 09:35", "Priya Raman", "Good idea, I'll raise it with Nina on Monday"),
    ("25/10/2026, 09:37", "Aisha Rahman", "Also thank Dev for the fast triage at 2am"),
    ("25/10/2026, 09:38", "Priya Raman", "Will do. Closing this chat thread for the weekend"),
    ("25/10/2026, 09:39", "Devansh Patel", "Thanks all"),
]


def make_chat(path: Path) -> None:
    lines = ["24/10/2026, 02:13 - Messages and calls are end-to-end encrypted. No one outside of "
             "this chat, not even WhatsApp, can read or listen to them.",
             "24/10/2026, 02:13 - Priya Raman created group \"Ops on-call\""]
    lines += [f"{t} - {who}: {msg}" for t, who, msg in CHAT_LINES]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


# ------------------------------------------------------------------------------------- CSV
def make_csv(path: Path) -> None:
    rows = [
        ("month", "service", "vendor", "severity", "tickets", "sla_target", "breaches",
         "sla_credit_sgd", "vendor_contact", "company_owner", "notes"),
        ("2026-07", "KYC OCR", "Acme Vendor", "Sev 1", "1", "4h response", "0", "0", "ken.ext", "hana.y", "Within SLA"),
        ("2026-07", "KYC OCR", "Acme Vendor", "Sev 2", "6", "next business day", "0", "0", "ken.ext", "hana.y", "Within SLA"),
        ("2026-08", "KYC OCR", "Acme Vendor", "Sev 1", "2", "4h response", "0", "0", "ken.ext", "hana.y", "Within SLA"),
        ("2026-08", "KYC OCR", "Acme Vendor", "Sev 2", "4", "next business day", "1", "0", "ken.ext", "hana.y",
         "One Sev 2 answered after 30h"),
        ("2026-09", "KYC OCR", "Acme Vendor", "Sev 1", "1", "4h response", "0", "0", "ken.ext", "leo.ops", "Within SLA"),
        ("2026-09", "KYC OCR", "Acme Vendor", "Sev 2", "5", "next business day", "0", "0", "ken.ext", "leo.ops", "Within SLA"),
        ("2026-10", "KYC OCR", "Acme Vendor", "Sev 1", "3", "4h response", "2", "1800", "ken.ext", "leo.ops",
         "October: 2 Sev 1 SLA breaches on the KYC OCR service (passport MRZ outage); "
         "SLA credit of SGD 1,800 owed by Acme Vendor"),
        ("2026-10", "KYC OCR", "Acme Vendor", "Sev 2", "7", "next business day", "0", "0", "ken.ext", "leo.ops", "Within SLA"),
    ]
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


# -------------------------------------------------------------------------------- Markdown
MD_TEXT = """# Onboarding buddy guide

*Company A · People & Workplace · owner: Hana Yusuf · v1, October 2026*

Every new joiner at Company A is paired with a buddy from another team for their first 30 days.

The buddy is not the new joiner's manager: they are the person to ask "where is..." and "who owns...".

## What a buddy does

1. Day 1: lunch together, a tour of the office and of the main Slack channels (#general, #random, #incidents).
2. Week 1: a 30-minute walk through the Company policies folder in the Company drive.
3. Weeks 2 to 4: a 15-minute check-in every Friday.
4. Day 30: a short note to the joiner's manager on how onboarding went.

## Buddy coffee allowance

Buddies can claim a buddy coffee allowance of SGD 40 per month for the 30 days, through the normal expense form.

Keep the receipt; no pre-approval needed.

## How to volunteer

Message Hana Yusuf by the 25th of the month. You need at least six months at Company A.

---
_Synthetic demonstration content for Walnut DataBridge; every name is invented._
"""


FIXED_ISO = "2026-10-30T00:00:00Z"
FIXED_DATE = (2026, 10, 30, 0, 0, 0)


def normalize_zip(path: Path) -> None:
    """Make an Office file byte-identical on every run: fixed core-property dates, sorted
    entries, fixed zip timestamps and permissions."""
    with zipfile.ZipFile(path) as zf:
        items = {n: zf.read(n) for n in zf.namelist()}
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for name in sorted(items):
            data = items[name]
            if name == "docProps/core.xml":
                t = data.decode("utf-8")
                t = re.sub(r"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)", r"\g<1>" + FIXED_ISO + r"\g<2>", t)
                data = t.encode("utf-8")
            zi = zipfile.ZipInfo(name, FIXED_DATE)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            out.writestr(zi, data)
    tmp.replace(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    make_docx(OUT / DOCX)
    make_xlsx(OUT / XLSX)
    make_pptx(OUT / PPTX)
    make_chat(OUT / CHAT)
    make_csv(OUT / CSVF)
    (OUT / MD).write_text(MD_TEXT, encoding="utf-8")
    for f in (DOCX, XLSX, PPTX):
        normalize_zip(OUT / f)
    PUBLIC.mkdir(parents=True, exist_ok=True)
    for f in ALL:
        shutil.copy2(OUT / f, PUBLIC / f)
        print(f"{(OUT / f).stat().st_size:>8} B  {f}")
    print(f"copied to {PUBLIC}")


if __name__ == "__main__":
    main()
